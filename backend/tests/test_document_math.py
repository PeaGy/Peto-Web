"""Công thức trong tài liệu Peto tạo: LaTeX → công thức Word thật (OMML) và hình vẽ trong PDF."""
import re
from io import BytesIO
from zipfile import ZipFile

import pytest
from docx import Document
from pypdf import PdfReader

from features.documents.export import MAX_TABLE_COLUMNS, parse_blocks, render_docx, render_pdf
from features.documents.math import markdown as math_markdown
from features.documents.math.latex import MathError, parse
from features.documents.math.layout import TextFonts, layout
from features.documents.math.omml_in import math_paragraph, omml_to_latex
from features.documents.math.omml_out import to_omml
from features.documents.word_reader import parse as parse_xml

NAMESPACES = ('xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math" '
              'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"')
M = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"

FORMULAS = [
    r"F = xyz + wxz + \bar{x}y\bar{z} + \overline{x}\,\overline{y}\,\overline{z}",
    r"\frac{x+1}{x^2-1}",
    r"\sqrt[3]{x+1} + \sqrt{2}",
    r"\sum_{i=1}^{n} i^2 = \frac{n(n+1)(2n+1)}{6}",
    r"\int_0^1 x\,dx = \frac{1}{2}",
    r"\lim_{x \to 0} \frac{\sin x}{x} = 1",
    r"\begin{cases} x + y = 3 \\ x - y = 1 \end{cases}",
    r"\begin{pmatrix} 1 & 2 \\ 3 & 4 \end{pmatrix}",
    r"\left( \frac{a}{b} \right)^2",
    r"f'(x) = 2x",
    r"\widehat{ABC} = 90^\circ",
    r"x \in \mathbb{R}",
    r"\text{nếu } x > 0",
    r"\binom{n}{k}",
    r"\log_2 x + \sin^2 x",
    r"\vec{AB} \perp \vec{CD}",
    r"\tg x \cdot \cotg x = 1",
    r"\Leftrightarrow x = 2 \lor x = -2",
    r"a &= b + c \\ &= d",
    r"\left\{ x \middle| x > 0 \right\}",
    r"\overbrace{1+\dots+1}^{n}",
    r"\xrightarrow{t} \quad \boxed{x=1} \quad \not\in",
]


def omml(source: str, display: bool = True):
    return parse_xml(f"<w:p {NAMESPACES}>{to_omml(parse(source, display))}</w:p>".encode())


def back_to_latex(paragraph) -> str:
    out = []
    for child in paragraph:
        if child.tag == M + "oMathPara":
            out.extend(math_paragraph(child))
        elif child.tag == M + "oMath":
            out.append(omml_to_latex(child))
    return " ".join(out)


@pytest.mark.parametrize("source", FORMULAS)
def test_formulas_become_word_equations_that_read_back_the_same(source):
    first = back_to_latex(omml(source))
    assert first
    # Đọc lại công thức Word vừa dựng rồi dựng tiếp phải ra cùng công thức: hai chiều khớp nhau (khoảng trắng của
    # LaTeX không mang nghĩa nên không so).
    assert back_to_latex(omml(first)).replace(" ", "") == first.replace(" ", "")


def test_word_equation_structures_follow_the_schema_order():
    paragraph = omml(r"\sum_{i=1}^{n} i \cdot \frac{a}{b} + \sqrt{x} + \overline{xy} + \begin{bmatrix} 1 \end{bmatrix}"
                     r" + \sin x + \hat{a}")
    nary = next(paragraph.iter(M + "nary"))
    assert [child.tag.replace(M, "") for child in nary] == ["naryPr", "sub", "sup", "e"]
    assert nary.find(f"{M}naryPr/{M}chr").get(M + "val") == "∑"
    assert next(paragraph.iter(M + "rad")).find(f"{M}radPr/{M}degHide") is not None
    assert next(paragraph.iter(M + "bar")).find(f"{M}barPr/{M}pos").get(M + "val") == "top"
    bracket = next(node for node in paragraph.iter(M + "d") if node.find(f"{M}e/{M}m") is not None)
    assert bracket.find(f"{M}dPr/{M}begChr").get(M + "val") == "["
    function = next(paragraph.iter(M + "func"))
    assert function.find(f"{M}fName/{M}r/{M}rPr/{M}sty").get(M + "val") == "p"
    assert next(paragraph.iter(M + "acc")).find(f"{M}accPr/{M}chr").get(M + "val") == "\u0302"
    # m:nor (chữ thường) không đi cùng m:sty: đậm/nghiêng của \textbf nằm ở w:rPr.
    text = omml(r"\textbf{đậm}")
    properties = next(text.iter(M + "rPr"))
    assert [child.tag.replace(M, "") for child in properties] == ["nor"]


@pytest.mark.parametrize("source,message", [
    (r"\frac{a}{", "Thiếu dấu }"),
    (r"x^2^3", "Hai chỉ số trên"),
    (r"\foo x", r"\foo chưa hỗ trợ"),
    (r"\left( x", r"\left thiếu \right"),
    (r"a & b", "Dấu & chỉ dùng"),
    (r"\begin{foo} \end{foo}", "chưa hỗ trợ"),
    ("}", "Thừa dấu }"),
    ("", "trống"),
])
def test_bad_formulas_are_refused_in_vietnamese(source, message):
    with pytest.raises(MathError, match=re.escape(message)):
        parse(source)


@pytest.mark.parametrize("text,formulas", [
    ("Giá $5 và $10.", []),
    ("Ta có $x^2$ và $$y=1$$.", [("x^2", False), ("y=1", True)]),
    ("Mã `$a$` giữ nguyên.", []),
    (r"\(a+b\) và \[c\]", [("a+b", False), ("c", True)]),
    ("$$\nF = a\n+ b\n$$\nnên", [("F = a\n+ b", True)]),
    ("```\n$x$\n```\n$y$", [("y", False)]),
    (r"Thoát \$5 không phải công thức.", []),
    ("$ x$ không mở, $x $ không đóng.", []),
])
def test_math_is_extracted_before_markdown(text, formulas):
    assert math_markdown.extract(text)[1] == formulas


def test_blocks_keep_formulas_in_headings_tables_and_display_lines():
    blocks = parse_blocks("# Bài $1$\n\nCho $F$ và\n\n$$\nF = a\n+ b\n$$\n\n| a | $|x|$ |\n|---|---|\n| 1 | 2 |\n")
    assert [block.kind for block in blocks] == ["heading", "paragraph", "math", "table"]
    assert blocks[0].runs[-1].math and blocks[0].runs[-1].text == "1"
    assert blocks[2].math == "F = a\n+ b"
    assert blocks[3].rows[0][1][0].math and blocks[3].rows[0][1][0].text == "|x|"


def test_bad_formula_is_reported_with_its_text():
    with pytest.raises(ValueError, match=r'Công thức "\\frac\{a\}\{" chưa dựng được'):
        parse_blocks("Rút gọn $\\frac{a}{$ nhé.")


def test_docx_has_real_word_equations_and_no_raw_latex():
    content = ("# Lời giải\n\nCho $F = \\bar{x}y$.\n\n$$\\int_0^1 x\\,dx = \\frac{1}{2}$$\n\n"
               "| $x$ | $F$ |\n|---|---|\n| 0 | 1 |\n")
    data = render_docx("Lời giải", content, "essay")
    with ZipFile(BytesIO(data)) as package:
        xml = package.read("word/document.xml").decode()
    assert xml.count("<m:oMath>") + xml.count("<m:oMath ") >= 4
    assert len(re.findall(r"<m:oMathPara[ >]", xml)) == 1
    assert "\\frac" not in xml and "$" not in xml
    document = Document(BytesIO(data))
    cell = document.tables[0].rows[1].cells[0].paragraphs[0]
    assert cell.paragraph_format.first_line_indent == 0
    # Đầu trang ghi tên tài liệu, không còn chữ "Nghị luận xã hội" cố định.
    assert document.sections[0].header.paragraphs[0].text == "Lời giải"


def test_pdf_draws_formulas_and_typed_math_symbols():
    content = ("Cho $F = \\bar{x}y + \\frac{a}{b}$ trong dòng.\n\n$$\\sum_{i=1}^{n} i = \\frac{n(n+1)}{2} \\tag{1}$$\n\n"
               "Ký hiệu gõ thẳng: x ≤ 5, A ⊂ B.\n")
    data = render_pdf("Lời giải", content, "essay")
    reader = PdfReader(BytesIO(data))
    fonts = "".join(str(font.get_object().get("/BaseFont")) for font in reader.pages[0]["/Resources"]["/Font"].values())
    assert "STIXTwoMath" in fonts
    text = reader.pages[0].extract_text()
    assert "(1)" in text and "≤" in text


@pytest.mark.parametrize("source,expected", [
    (r"~p \lor q", r"{\sim}p \lor q"), (r"(~p)", r"({\sim}p)"), (r"~~p", r"{\sim}{\sim}p"),
    (r"p \to ~q", r"p \to {\sim}q"), ("a~b", "a~b"), (r"\text{x}~y", r"\text{x}~y"),
])
def test_tilde_before_an_operand_is_negation_like_the_chat(source, expected):
    from features.documents.math.latex import tilde_negation
    assert tilde_negation(source) == expected


def test_layout_sizes_follow_display_and_inline_rules():
    fonts = TextFonts("Helvetica", "Helvetica-Bold", "Helvetica-Oblique")
    inline = layout(parse(r"\sum_{i=1}^{n} \frac{a}{b}"), 12, fonts)
    display = layout(parse(r"\sum_{i=1}^{n} \frac{a}{b}", True), 12, fonts)
    assert display.height > inline.height and display.depth > inline.depth
    assert inline.width > 0 and inline.height > 0


def test_wide_truth_tables_fit_but_not_unbounded():
    header = "| " + " | ".join(f"c{n}" for n in range(9)) + " |\n|" + "---|" * 9 + "\n"
    assert parse_blocks(header)[0].kind == "table"
    too_wide = "| " + " | ".join("c" for _ in range(MAX_TABLE_COLUMNS + 1)) + " |\n|" + "---|" * (MAX_TABLE_COLUMNS + 1)
    with pytest.raises(ValueError, match=f"{MAX_TABLE_COLUMNS} cột"):
        parse_blocks(too_wide)
