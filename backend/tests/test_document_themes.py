"""Kiểu trình bày tài liệu (Khung đôi, Dải màu, Tối giản, nghị luận) và trang bìa ghi ở đầu nội dung."""
import re
from io import BytesIO
from zipfile import ZipFile

import pypdfium2 as pdfium
import pytest
from docx import Document
from docx.shared import Cm

from conftest import read_events
from features.documents.cover import DOTS, split_cover
from features.documents.export import column_widths, parse_blocks, render_docx, render_pdf
from features.documents.jobs import build_files
from features.documents.tools import SCHEMA

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

COVER = """---
trường: Trường Đại học Mẫu
khoa: Khoa Công nghệ Thông tin
loại: Báo cáo thực hành giữa kỳ
môn: Dịch vụ mạng
đề tài: Windows Defender Application Control
giảng viên: ThS. Nguyễn Văn X
nhóm: Nhóm 7
thành viên:
- Nguyễn Văn A | 124000001 | Nhóm trưởng
- Trần Thị B - 124000002
nơi: TP. Hồ Chí Minh
---
"""
BODY = """# Báo cáo thực hành

[TOC]

## I. Giới thiệu

Đoạn mở đầu.

## II. Thực hiện

Bảng: So sánh công cụ

| Công cụ | Có từ | Lab |
|---|---|---|
| AppLocker | Windows 7 | Không |

```
\\\\172.16.64.1\\DungChung
secpol.msc
```

![Cửa sổ Local Security Settings](khung-anh)
"""


def xml_of(data: bytes, part: str = "word/document.xml") -> str:
    with ZipFile(BytesIO(data)) as package:
        return package.read(part).decode()


def all_parts(data: bytes) -> str:
    with ZipFile(BytesIO(data)) as package:
        return "".join(package.read(name).decode() for name in package.namelist()
                       if name.startswith("word/") and name.endswith(".xml"))


def test_cover_front_matter_is_read_with_aliases_and_members():
    cover, rest = split_cover(COVER + BODY)
    assert cover.school == "Trường Đại học Mẫu" and cover.faculty == "Khoa Công nghệ Thông tin"
    assert cover.kind == "Báo cáo thực hành giữa kỳ" and cover.course == "Dịch vụ mạng"
    assert [(member.name, member.code, member.note) for member in cover.members] == [
        ("Nguyễn Văn A", "124000001", "Nhóm trưởng"), ("Trần Thị B", "124000002", "")]
    assert cover.when() == "TP. Hồ Chí Minh, " + cover.when().split(", ", 1)[1]
    assert rest.lstrip().startswith("# Báo cáo thực hành")
    unaccented, _ = split_cover("---\ntruong: A\nGVHD: B\nsinh vien: C; D\nlogo: anh-2\n---\nThân bài")
    assert (unaccented.school, unaccented.instructor, unaccented.logo) == ("A", "B", 2)
    assert [member.name for member in unaccented.members] == ["C", "D"]


@pytest.mark.parametrize("content", ["Không có bìa.\n\n---\n\nĐoạn sau đường kẻ.", "---\ntrường: A\nkhông đóng"])
def test_content_without_cover_block_is_left_alone(content):
    cover, rest = split_cover(content)
    assert cover is None and rest == content


def test_unknown_cover_key_is_refused_with_the_allowed_keys():
    with pytest.raises(ValueError, match="khóa: trường, khoa"):
        parse_blocks("---\nhiệu trưởng: Ai đó\n---\nNội dung")


def test_missing_cover_information_stays_blank_and_date_is_filled():
    cover, _ = split_cover("---\nmôn: Mạng máy tính\n---\n")
    text = " ".join(paragraph.text for table in Document(BytesIO(render_docx("Bài", "---\nmôn: Mạng máy tính\n---\nThân.", "classic"))).tables
                    for row in table.rows for cell in row.cells for paragraph in cell.paragraphs)
    assert f"TRƯỜNG {DOTS}" in text and f"KHOA {DOTS}" in text and DOTS in text.split("Giảng viên hướng dẫn: ")[1]
    assert re.search(r"Tháng \d+ năm 20\d\d", cover.when())


@pytest.mark.parametrize("style", ["classic", "band", "minimal"])
def test_word_reports_follow_the_vietnamese_standard_with_a_cover(style):
    data = render_docx("Báo cáo thực hành", COVER + BODY, style)
    document = Document(BytesIO(data))
    section = document.sections[0]
    assert abs(section.page_width - Cm(21)) < Cm(0.01) and abs(section.page_height - Cm(29.7)) < Cm(0.01)
    assert abs(section.left_margin - Cm(3)) < Cm(0.01) and abs(section.right_margin - Cm(2)) < Cm(0.01)
    assert section.different_first_page_header_footer, "bìa không có đầu, chân trang"
    normal = document.styles["Normal"]
    assert normal.font.name == "Times New Roman" and normal.font.size.pt == 13
    assert normal.paragraph_format.line_spacing == 1.5
    xml = xml_of(data)
    # Bảng thành viên nằm lồng trong bảng ẩn viền của bìa: tìm trong toàn bộ chữ của tệp.
    cover_text = " ".join(re.findall(r"<w:t(?: [^>]*)?>([^<]*)</w:t>", xml))
    for expected in ["Trường Đại học Mẫu", "Windows Defender Application Control", "Nguyễn Văn A", "124000002",
                     "ThS. Nguyễn Văn X"]:
        assert expected in cover_text, expected
    if style == "classic":
        assert 'w:display="firstPage"' in xml and "thickThinSmallGap" in xml
    if style == "band":
        assert "wps:wsp" in xml and 'relativeFrom="page"' in xml and 'behindDoc="1"' in xml
        header = xml_of(data, "word/header1.xml") if "word/header1.xml" in ZipFile(BytesIO(data)).namelist() else all_parts(data)
        assert "Dịch vụ mạng · Nhóm 7" in header
    if style == "minimal":
        assert "<w:smallCaps/>" in xml
    texts = [paragraph.text for paragraph in document.paragraphs]
    # Chú thích bảng đánh số và không còn đoạn "Bảng:" gốc; khung chừa ảnh có chú thích Hình 1.
    separator = ":" if style == "band" else "."
    assert f"Bảng 1{separator} So sánh công cụ" in texts and "Bảng: So sánh công cụ" not in texts
    assert f"Hình 1{separator} Cửa sổ Local Security Settings" in texts
    assert 'w:val="dashed"' in xml
    code = [paragraph for paragraph in document.paragraphs if paragraph.style.name == "Code"]
    assert [paragraph.text for paragraph in code] == ["\\\\172.16.64.1\\DungChung", "secpol.msc"]
    # Tiêu đề trùng tên tài liệu không lặp lại sau bìa.
    assert "Báo cáo thực hành" not in texts


def test_classic_headings_are_caps_by_formatting_and_toc_has_dotted_tabs():
    data = render_docx("Báo cáo thực hành", COVER + BODY, "classic")
    document = Document(BytesIO(data))
    assert document.styles["Heading 1"].font.all_caps
    heading = next(paragraph for paragraph in document.paragraphs if paragraph.style.name == "Heading 1")
    assert heading.text == "I. Giới thiệu", "chữ gốc giữ nguyên, Word hiện in hoa"
    styles = xml_of(data, "word/styles.xml")
    assert re.search(r'w:styleId="TOC1".*?w:leader="dot"', styles, re.S)


def test_build_fills_word_contents_with_pdf_page_numbers():
    files = build_files("Báo cáo thực hành", COVER + BODY, "band", "docx")
    document = Document(BytesIO(files["docx"]))
    entries = [paragraph.text for paragraph in document.paragraphs if paragraph.style.name.startswith("toc")]
    assert entries == ["I. Giới thiệu\t3", "II. Thực hiện\t3"]
    assert files["pages"] == 3


@pytest.mark.parametrize("style", ["classic", "band", "minimal", "essay"])
def test_pdf_draws_the_cover_first_and_the_text_after(style):
    data = render_pdf("Báo cáo thực hành", COVER + BODY, style)
    pdf = pdfium.PdfDocument(data)
    first = " ".join(pdf[0].get_textpage().get_text_range().casefold().split())
    assert "windows defender application control" in first and "nguyễn văn a" in first
    second = " ".join(pdf[1].get_textpage().get_text_range().casefold().split())
    assert "mục lục" in second and "giới thiệu" in second
    if style == "band":
        assert "dịch vụ mạng · nhóm 7" in second and "trang 2" in second


def test_without_cover_the_opening_heading_is_the_title_and_page_one_has_no_header():
    content = "# Lời giải bài tập 5: Đại số Bool\n\n## Bài 1\n\nCho $F = x$.\n\n## Bài 2\n\nXong."
    document = Document(BytesIO(render_docx("Lời giải bài tập 5", content, "minimal")))
    texts = [paragraph.text for paragraph in document.paragraphs]
    assert texts[0] == "Lời giải bài tập 5: Đại số Bool" and texts.count("Lời giải bài tập 5: Đại số Bool") == 1
    assert document.paragraphs[1].style.name == "Heading 1"
    section = document.sections[0]
    assert section.different_first_page_header_footer
    assert "PAGE" in section.first_page_footer._element.xml and not section.first_page_header.paragraphs[0].text
    page = pdfium.PdfDocument(render_pdf("Lời giải bài tập 5", content, "minimal"))[0].get_textpage().get_text_range()
    assert page.count("Lời giải bài tập 5") == 1


def test_old_report_style_is_rendered_as_classic():
    document = Document(BytesIO(render_docx("Kế hoạch", "# Kế hoạch\n\nNội dung.", "report")))
    assert document.styles["Normal"].font.name == "Times New Roman"
    assert abs(document.sections[0].page_width - Cm(21)) < Cm(0.01)


def test_table_columns_fit_their_longest_word():
    from features.documents.export import Run
    rows = [[[Run("Công cụ")], [Run("Cách nhận file theo chữ ký số, hash và nhà phát hành")]],
            [[Run("AppLocker")], [Run("Đường dẫn")]]]
    widths = column_widths(rows, 453)
    assert abs(sum(widths) - 453) < 0.01
    assert widths[0] >= len("AppLocker") * 6.3 + 12
    # Công thức tính theo ký hiệu nhìn thấy: cột \bar{x}y không rộng gấp mấy cột x.
    math = [[[Run("x", math=True)], [Run("\\bar{x}\\bar{y}\\bar{z}", math=True)]]]
    narrow = column_widths(math, 100)
    assert narrow[1] < narrow[0] * 2


FILLER = "Đoạn văn mẫu dùng để đẩy nội dung xuống gần cuối trang, đủ dài để chiếm vài dòng khi căn đều hai bên lề. " * 2
ROWS = "| A | B |\n|---|---|\n"
FOLLOWERS = {
    "bảng ngắn có chú thích": ("Bảng: Bảng ngắn\n\n" + ROWS + "".join(f"| {i} | x |\n" for i in range(5)), "bảng 1", False),
    "bảng dài": (ROWS + "".join(f"| {i} | dòng {i} |\n" for i in range(40)), "dòng 0", True),
    "bảng dài có chú thích": ("Bảng: Bảng dài\n\n" + ROWS + "".join(f"| {i} | dòng {i} |\n" for i in range(40)), "bảng 1",
                              False),
    "đoạn dài": ("Đoạn dài sau đề mục. " * 60, "đoạn dài", True),
}


@pytest.mark.parametrize("follower", list(FOLLOWERS))
def test_pdf_headings_stay_with_the_start_of_the_next_block(follower):
    """Đề mục ở cuối trang: không nằm lại một mình (khối sau là bảng ngắn, ảnh), và không kéo cả đoạn văn hay bảng dài
    sang trang sau để trống nửa trang như KeepTogether của ReportLab; Word cũng chỉ giữ đề mục với phần đầu khối sau."""
    after, marker, splittable = FOLLOWERS[follower]
    for count in range(6, 14):
        content = "\n\n".join([FILLER] * count) + "\n\n## Đề mục thử\n\n" + after + "\n\n## Kết\n\nXong."
        pdf = pdfium.PdfDocument(render_pdf("Thử", content, "minimal"))
        texts = [pdf[index].get_textpage().get_text_range().casefold() for index in range(len(pdf))]
        page = next(index for index, text in enumerate(texts) if "đề mục thử" in text)
        assert marker in texts[page].split("đề mục thử", 1)[1], f"{count} đoạn: đề mục nằm lại cuối trang {page + 1}"
        if splittable and page:
            text = pdf[page - 1].get_textpage()
            # Đáy dòng chữ thấp nhất trên chân trang (số trang nằm dưới 45 pt), so với lề dưới 2 cm.
            bottom = min(box[1] for box in map(text.get_charbox, range(text.count_chars())) if box[1] > 45)
            assert bottom - 2 * 28.3465 < 110, f"{count} đoạn: trang {page} để trống {bottom:.0f} pt ở cuối"


def test_word_keeps_short_tables_on_one_page_and_slots_with_their_caption():
    long_rows = "".join(f"| {i} | dòng {i} |\n" for i in range(20))
    content = f"Bảng: Ngắn\n\n{ROWS}| 1 | x |\n| 2 | y |\n\nĐoạn.\n\n{ROWS}{long_rows}\n![Màn hình](khung-anh)"
    document = Document(BytesIO(render_docx("Bảng", content, "classic")))
    short, long = document.tables[0], document.tables[1]
    keeps = [row.cells[0].paragraphs[0].paragraph_format.keep_with_next for row in short.rows]
    assert keeps == [True, True, None], "mọi hàng trừ hàng cuối giữ với hàng sau"
    assert all(row._tr.trPr.find(f"{W}cantSplit") is not None for row in short.rows)
    assert not any(row.cells[0].paragraphs[0].paragraph_format.keep_with_next for row in long.rows), "bảng dài tách trang"
    slot = document.tables[2].rows[0].cells[0].paragraphs[0]
    assert slot.paragraph_format.keep_with_next and "Chỗ dán ảnh" in slot.text


def test_tool_offers_the_three_styles_and_the_essay():
    style = SCHEMA["parameters"]["properties"]["style"]
    assert style["enum"] == ["classic", "band", "minimal", "essay"]
    content = SCHEMA["parameters"]["properties"]["content"]["description"]
    assert "khung-anh" in content and "Bảng: tên bảng" in content and "thành viên" in content


async def test_hand_edits_keep_the_chosen_style(client):
    events = await read_events(await client.post("/api/chat", json={"message": "__baocao__:band"}))
    artifact = next(event["artifact"] for event in events if event["type"] == "artifact")
    assert artifact["style"] == "band"
    path = f"/api/documents/{artifact['id']}"
    current = (await client.get(path)).json()
    revised = await client.post(path + "/versions", json={"base_version": current["version"], "title": current["title"],
                                                         "content": current["content"] + "\n\nThêm một đoạn.",
                                                         "style": current["style"]})
    assert revised.status_code == 200 and revised.json()["style"] == "band"
    word = (await client.get(path + "/export/docx", params={"version": revised.json()["version"]})).content
    assert "wps:wsp" in xml_of(word), "bản sửa tay vẫn có bìa Dải màu"
    entries = [paragraph.text for paragraph in Document(BytesIO(word)).paragraphs if paragraph.style.name.startswith("toc")]
    assert entries and all(re.search(r"\t\d+$", entry) for entry in entries), "mục lục bản sửa tay cũng có số trang"
