from io import BytesIO
from zipfile import ZipFile
import asyncio
import pytest
from pypdf import PdfReader
from docx import Document
from conftest import TEST_OWNER, read_events
import db
import document_api
from document_export import parse_blocks, render_docx, render_pdf

CONTENT = '''# Kế hoạch học tập

Tiếng Việt có dấu: Nguyễn, trường học, kỹ thuật. **Đậm** và *nghiêng*.

## Công việc

1. Đọc tài liệu
2. Thực hành

| Ngày | Việc cần làm |
| --- | --- |
| Thứ hai | Ôn tập |
| Thứ tư | Viết báo cáo |
'''


async def create(client):
    conversation = await db.create_conversation(TEST_OWNER, 'Tài liệu')
    response = await client.post('/api/documents', json={'conversation_id': conversation, 'title': 'Kế hoạch học tập', 'content': CONTENT})
    assert response.status_code == 200
    return conversation, response.json()


async def test_ownership_and_authentication(client, anon_client):
    conversation, document = await create(client)
    other = await db.create_conversation('other-owner', 'Không được đọc')
    assert (await anon_client.get(f"/api/documents/{document['id']}")).status_code == 401
    assert (await client.get('/api/documents', params={'conversation_id': other})).status_code == 404
    assert (await client.post('/api/documents', json={'conversation_id': other, 'title': 'X', 'content': 'X'})).status_code == 404
    await db.delete_conversation(TEST_OWNER, conversation)
    assert (await client.get(f"/api/documents/{document['id']}")).status_code == 404


async def test_versions_export_exact_revision_and_conflict(client):
    conversation, document = await create(client)
    path = f"/api/documents/{document['id']}"
    payload = {'base_version': 1, 'title': 'Bản hai', 'content': 'Nội dung đã sửa'}
    responses = await asyncio.gather(client.post(path + '/versions', json=payload), client.post(path + '/versions', json=payload))
    assert sorted(response.status_code for response in responses) == [200, 409]
    latest = (await client.get(path)).json()
    assert latest['version'] == 2
    assert [v['version'] for v in latest['versions']] == [2, 1]
    assert (await client.get(path, params={'version': 1})).json()['content'] == CONTENT.strip()
    for format in ['docx', 'pdf']:
        response = await client.get(path + '/export/' + format, params={'version': 1})
        assert response.status_code == 200, response.text[:200] if response.status_code != 200 else ''
        assert 'attachment;' in response.headers['content-disposition']
        assert response.headers['cache-control'] == 'private, no-store'
        if format == 'docx': assert Document(BytesIO(response.content)).paragraphs[0].text == 'Kế hoạch học tập'
        else: assert 'Kế hoạch học tập' in PdfReader(BytesIO(response.content)).pages[0].extract_text()
    assert (await client.get('/api/documents', params={'conversation_id': conversation})).json()['documents'][0]['version'] == 2
    assert (await client.delete(path)).status_code == 200
    assert (await client.get(path + '/export/pdf', params={'version': 1})).status_code == 404


async def test_other_owner_cannot_read_revise_export_or_delete(client):
    from document_store import save_document
    conversation = await db.create_conversation('other-owner', 'Private')
    document = await save_document('other-owner', conversation, 'Private', 'Secret')
    path = f"/api/documents/{document['id']}"
    assert (await client.get(path)).status_code == 404
    assert (await client.post(path + '/versions', json={'base_version': 1, 'title': 'X', 'content': 'X'})).status_code == 404
    assert (await client.get(path + '/export/pdf', params={'version': 1})).status_code == 404
    assert (await client.delete(path)).status_code == 404


async def test_input_limits_and_busy_export(client):
    conversation, document = await create(client)
    assert (await client.post('/api/documents', json={'conversation_id': conversation, 'title': ' ', 'content': 'abc'})).status_code == 400
    assert (await client.post('/api/documents', json={'conversation_id': conversation, 'title': 'X', 'content': 'a' * 60001})).status_code == 422
    await document_api._render_lock.acquire()
    try:
        assert (await client.get(f"/api/documents/{document['id']}/export/pdf?version=1")).status_code == 429
    finally: document_api._render_lock.release()


def test_real_files_preserve_accents_structure_and_escape_markup():
    content = CONTENT + '\n\n<script>alert(1)</script> & text\n\n![private](http://127.0.0.1/secret)\n\n[Nguồn](https://example.org/reference)\n'
    word = render_docx('Kế hoạch học tập', content)
    document = Document(BytesIO(word))
    assert len(document.tables) == 1
    assert document.tables[0].cell(1, 0).text == 'Thứ hai'
    assert sum(p.text == 'Kế hoạch học tập' for p in document.paragraphs) == 1
    assert any(run.bold for paragraph in document.paragraphs for run in paragraph.runs)
    with ZipFile(BytesIO(word)) as package:
        assert not any('media/' in name or 'vba' in name for name in package.namelist())
    pdf = PdfReader(BytesIO(render_pdf('Kế hoạch học tập', content)))
    text = '\n'.join(page.extract_text() for page in pdf.pages)
    for expected in ['Nguyễn', 'kỹ thuật', 'Thứ hai', '<script>alert(1)</script>', '[Ảnh: private]', 'https://example.org/reference']:
        assert expected in text
    assert '/OpenAction' not in pdf.trailer['/Root']


def test_long_table_paginates_and_repeats_header():
    content = '| Ngày | Nội dung |\n| --- | --- |\n' + '\n'.join(f'| Ngày {i} | ' + 'Học tập đều đặn. ' * 15 + '|' for i in range(50))
    pdf = PdfReader(BytesIO(render_pdf('Báo cáo', content)))
    assert len(pdf.pages) > 1
    assert all('Nội dung' in page.extract_text() for page in pdf.pages)
    assert 'Ngày 49' in pdf.pages[-1].extract_text()
    with pytest.raises(ValueError): parse_blocks('| a |\n|---|\n' + '| b |\n' * 101)


async def test_document_mode_goes_through_normal_chat_and_saves_reply(client):
    response = await client.post('/api/chat', json={'message': 'Soạn kế hoạch học', 'document_mode': True})
    events = await read_events(response)
    reply = ''.join(event.get('text', '') for event in events if event['type'] == 'delta')
    assert reply.startswith('Đã tạo tệp mẫu')
    assert any(event['type'] == 'artifact' for event in events)
    assert events[-1]['type'] == 'done'
