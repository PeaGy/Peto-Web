"""Danh sách đánh số thật, ảnh của hội thoại và mục lục trong tài liệu Peto tạo; hàng chờ dựng tài liệu."""
import asyncio
import base64
import json
import re
from io import BytesIO
from zipfile import ZipFile

import pypdfium2 as pdfium
import pytest
from docx import Document
from PIL import Image
from pypdf import PdfReader

import storage as db
from features.chat import service as chat_service
from features.chat import titles
from ai.xai import build_input_payload
from conftest import TEST_OWNER, read_events
from features.documents.export import parse_blocks, render_docx, render_pdf
from features.documents.images import prepare
from features.documents.jobs import RenderBusy, RenderQueue
from features.documents.tools import DocumentSession

W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'

LISTS = '''# Kế hoạch

1. Đọc tài liệu
2. Thực hành
   - Viết code
   - Chạy thử
3. Nộp bài
   1. Bản nháp
   2. Bản cuối

Đoạn chen giữa.

5. Bắt đầu từ năm

   Đoạn thứ hai của mục năm.
6. Sáu
'''

CONTENTS = '''# Báo cáo

Mở đầu ngắn.

[TOC]

## Phần một

Nội dung một.

### Chi tiết

Nội dung chi tiết.

## Phần hai

Nội dung hai.
'''


def png(width=900, height=500, color=(40, 120, 200), mode='RGB'):
    image = Image.new(mode, (width, height), color if mode == 'RGB' else (*color, 120))
    output = BytesIO()
    image.save(output, 'PNG')
    return output.getvalue()


def attachment(name, data):
    return {'name': name, 'mime': 'image/png', 'data': base64.b64encode(data).decode()}


def numbering_of(paragraph):
    properties = paragraph._p.pPr.numPr if paragraph._p.pPr is not None else None
    return None if properties is None else (properties.ilvl.val, properties.numId.val)


async def chat(client, message, conversation=None, images=()):
    payload = {'message': message, 'attachments': list(images)}
    if conversation: payload['conversation_id'] = conversation
    events = await read_events(await client.post('/api/chat', json=payload))
    assert events[-1]['type'] == 'done', events[-1]
    return next(event['conversation_id'] for event in events if event['type'] == 'meta')


def test_lists_become_real_word_numbering_that_restarts_and_nests():
    document = Document(BytesIO(render_docx('Kế hoạch', LISTS)))
    paragraphs = {paragraph.text: paragraph for paragraph in document.paragraphs}
    numbered = {text: numbering_of(paragraphs[text]) for text in
                ['Đọc tài liệu', 'Thực hành', 'Viết code', 'Chạy thử', 'Nộp bài', 'Bản nháp', 'Bản cuối', 'Bắt đầu từ năm', 'Sáu']}
    assert all(numbered.values()), numbered
    # Không còn "1. " hay "• " gõ tay vào đầu đoạn: số do Word đánh.
    assert not any(re.match(r'(\d+\.|•|–) ', paragraph.text) for paragraph in document.paragraphs)
    first, nested_bullets, nested_steps, later = (numbered['Đọc tài liệu'][1], numbered['Viết code'][1],
                                                  numbered['Bản nháp'][1], numbered['Bắt đầu từ năm'][1])
    assert numbered['Thực hành'] == numbered['Nộp bài'] == (0, first)
    assert numbered['Chạy thử'] == (1, nested_bullets) and numbered['Bản cuối'] == (1, nested_steps)
    assert len({first, nested_bullets, nested_steps, later}) == 4, 'mỗi danh sách tự đánh số lại từ đầu'
    # Đoạn thứ hai của một mục thẳng hàng với chữ, không có số.
    continuation = paragraphs['Đoạn thứ hai của mục năm.']
    assert numbering_of(continuation) is None and continuation.paragraph_format.left_indent.pt == 36
    numbering = document.part.numbering_part.element

    def level(num_id, ilvl):
        abstract_id = next(num for num in numbering.findall(f'{W}num') if num.numId == num_id).abstractNumId.val
        abstract = next(item for item in numbering.findall(f'{W}abstractNum') if int(item.get(f'{W}abstractNumId')) == abstract_id)
        found = next(item for item in abstract.findall(f'{W}lvl') if item.get(f'{W}ilvl') == str(ilvl))
        return {child.tag.replace(W, ''): child.get(f'{W}val') for child in found if child.get(f'{W}val') is not None}

    assert level(first, 0) == {'start': '1', 'numFmt': 'decimal', 'lvlText': '%1.', 'lvlJc': 'left'}
    assert level(nested_bullets, 1)['numFmt'] == 'bullet' and level(nested_bullets, 1)['lvlText'] == '•'
    assert level(nested_steps, 1)['numFmt'] == 'lowerLetter'
    assert level(later, 0)['start'] == '5'

    text = PdfReader(BytesIO(render_pdf('Kế hoạch', LISTS))).pages[0].extract_text()
    for label, item in [('1.', 'Đọc tài liệu'), ('3.', 'Nộp bài'), ('•', 'Viết code'), ('a.', 'Bản nháp'), ('b.', 'Bản cuối'),
                        ('5.', 'Bắt đầu từ năm'), ('6.', 'Sáu')]:
        assert re.search(re.escape(label) + r'\s*' + item, text), (label, text)


def test_contents_marker_builds_word_field_and_pdf_page_numbers():
    word = render_docx('Báo cáo', CONTENTS)
    document = Document(BytesIO(word))
    texts = [paragraph.text for paragraph in document.paragraphs]
    assert '[TOC]' not in texts
    heading = texts.index('Mục lục')
    assert texts[heading + 1:heading + 4] == ['Phần một', 'Chi tiết', 'Phần hai']
    assert document.paragraphs[heading + 2].paragraph_format.left_indent.pt == 18
    body = document.element.body
    assert 'TOC \\o "1-3" \\h \\z \\u' in ''.join(item.text for item in body.iter(f'{W}instrText'))
    begin = next(item for item in body.iter(f'{W}fldChar') if item.get(f'{W}fldCharType') == 'begin')
    assert begin.get(f'{W}dirty') == 'true', 'Word dựng lại mục lục có số trang khi mở'
    assert any(item.get(f'{W}type') == 'page' for item in body.iter(f'{W}br'))

    data = render_pdf('Báo cáo', CONTENTS)
    pdf = PdfReader(BytesIO(data))
    # pdfium đọc chữ theo đúng thứ tự trên trang; pypdf đọc dòng chấm dẫn trước tên mục.
    first = pdfium.PdfDocument(data)[0].get_textpage().get_text_range()
    # Kiểu Khung đôi in hoa "MỤC LỤC" và đề mục cấp 1 trong bản PDF.
    assert 'Mở đầu ngắn.' in first and 'mục lục' in first.casefold()
    for entry in ['Phần một', 'Chi tiết', 'Phần hai']:
        assert re.search(entry.casefold() + r'[ .]*2\b', first.casefold()), first
        assert entry.casefold() in pdf.pages[1].extract_text().casefold()

    def titles_of(outline):
        return [titles_of(item) if isinstance(item, list) else item.title for item in outline]

    assert titles_of(pdf.outline) == ['Phần một', ['Chi tiết'], 'Phần hai']


def test_contents_marker_needs_headings_and_is_literal_elsewhere():
    for content in ['Chỉ có đoạn văn.\n\n[Mục lục]\n', '- [TOC]\n']:
        document = Document(BytesIO(render_docx('X', content)))
        assert 'Mục lục' not in [paragraph.text for paragraph in document.paragraphs]
        assert not list(document.element.body.iter(f'{W}instrText'))
    assert [block.kind for block in parse_blocks('[mục lục]\n\n## A')] == ['toc', 'heading']


def test_word_styles_use_the_chosen_fonts_instead_of_theme_fonts():
    # Mọi kiểu (và "report" cũ, dựng như Khung đôi) đều Times New Roman như các trường yêu cầu.
    for layout, font in [(name, 'Times New Roman') for name in ('classic', 'band', 'minimal', 'essay', 'report')]:
        with ZipFile(BytesIO(render_docx('Bài', '# Bài\n\n## Mở bài\n\nĐoạn.', layout))) as package:
            styles = package.read('word/styles.xml').decode()
        for style_id in ['Title', 'Heading1', 'Heading2', 'Caption']:
            block = re.search(r'<w:style [^>]*w:styleId="%s".*?</w:style>' % style_id, styles, re.S).group(0)
            assert 'Theme=' not in re.search(r'<w:rFonts[^>]*/>', block).group(0), (layout, style_id)
            assert f'w:ascii="{font}"' in block
        assert '<w:pBdr>' not in re.search(r'<w:style [^>]*w:styleId="Title".*?</w:style>', styles, re.S).group(0)


def test_prepare_shrinks_keeps_transparency_and_refuses_bombs():
    wide = prepare(png(3200, 800))
    assert (wide.width, wide.height) == (1600, 400) and wide.data.startswith(b'\x89PNG')
    assert prepare(png(mode='RGBA')).data.startswith(b'\x89PNG')
    photo = BytesIO()
    Image.effect_noise((800, 600), 60).convert('RGB').save(photo, 'WEBP')
    assert prepare(photo.getvalue()).data.startswith(b'\xff\xd8'), 'ảnh chụp nén JPEG'
    bomb = BytesIO()
    Image.new('1', (8000, 6000)).save(bomb, 'PNG')
    for data in [bomb.getvalue(), b'not an image']:
        with pytest.raises(ValueError):
            prepare(data)


async def test_conversation_images_go_into_both_files_with_captions(client):
    conversation = await chat(client, 'Ảnh đầu tiên', images=[attachment('bieu-do.png', png())])
    await chat(client, 'Ảnh thứ hai', conversation, [attachment('logo.png', png(400, 300, mode='RGBA'))])
    session = DocumentSession(TEST_OWNER, conversation)
    content = '# Báo cáo ảnh\n\nMở đầu.\n\n![Biểu đồ doanh thu](anh-1)\n\nGiữa hai ảnh. ![](anh-2) Sau ảnh.\n'
    result = await session.create(json.dumps({'title': 'Báo cáo ảnh', 'content': content, 'format': 'docx', 'style': 'classic'}))
    assert result['ok'], result
    base = f"/api/documents/{result['artifact']['id']}"
    word = (await client.get(base + '/export/docx?version=1')).content
    with ZipFile(BytesIO(word)) as package:
        assert len([name for name in package.namelist() if name.startswith('word/media/')]) == 2
    texts = [paragraph.text for paragraph in Document(BytesIO(word)).paragraphs]
    # Hình có chú thích được đánh số tự động; hình không chú thích thì không.
    assert 'Hình 1. Biểu đồ doanh thu' in texts and 'Giữa hai ảnh.' in texts and 'Sau ảnh.' in texts
    assert not any('anh-' in text or '[Ảnh' in text for text in texts)
    pdf = PdfReader(BytesIO((await client.get(base + '/export/pdf?version=1')).content))
    assert sum(len(page.images) for page in pdf.pages) == 2
    assert 'Biểu đồ doanh thu' in pdf.pages[0].extract_text()


async def test_image_numbers_are_checked_per_owner_and_conversation(client):
    conversation = await chat(client, 'Một ảnh', images=[attachment('a.png', png())])
    arguments = lambda content: json.dumps({'title': 'Ảnh', 'content': content, 'format': 'pdf', 'style': 'report'})
    missing = await DocumentSession(TEST_OWNER, conversation).create(arguments('![x](anh-3)'))
    assert 'Ảnh 3' in missing['error'] and 'Ảnh 1' in missing['error']
    empty = await db.create_conversation(TEST_OWNER, 'Không có ảnh')
    assert 'chưa có ảnh nào' in (await DocumentSession(TEST_OWNER, empty).create(arguments('![x](anh-1)')))['error']
    # Tài khoản khác biết mã hội thoại cũng không đọc được ảnh: truy vấn lọc theo chủ trong SQL.
    stranger = DocumentSession('other-owner', conversation)
    assert 'chưa có ảnh nào' in (await stranger.create(arguments('![x](anh-1)')))['error']
    assert not stranger.created


async def test_hand_edited_draft_shows_text_for_images_it_cannot_find(client):
    conversation = await chat(client, 'Ảnh bìa', images=[attachment('bia.png', png())])
    draft = (await client.post('/api/documents', json={'conversation_id': conversation, 'title': 'Bản tay',
                                                       'content': '![Ảnh bìa](anh-1)\n\n![Không có](anh-9)'})).json()
    word = (await client.get(f"/api/documents/{draft['id']}/export/docx?version=1")).content
    with ZipFile(BytesIO(word)) as package:
        assert len([name for name in package.namelist() if name.startswith('word/media/')]) == 1
    assert '[Ảnh 9: Không có]' in [paragraph.text for paragraph in Document(BytesIO(word)).paragraphs]


async def test_model_sees_each_image_with_its_number(client, monkeypatch):
    seen = []

    class Spy:
        async def stream(self, **kwargs):
            if titles.TITLE_MARKER in kwargs['system_prompt']:
                yield 'Tiêu đề'
                return
            seen.append(kwargs['messages'])
            yield 'Đã xem.'

    monkeypatch.setattr(chat_service, 'get_provider', lambda model='peto': Spy())
    conversation = await chat(client, 'Hai ảnh', images=[attachment('a.png', png()), attachment('b.png', png(color=(9, 9, 9)))])
    await chat(client, 'Thêm một ảnh', conversation, [attachment('c.png', png(color=(200, 0, 0)))])
    assert sorted((item.name, item.number) for message in seen[-1] for item in message.attachments) == [
        ('a.png', 1), ('b.png', 2), ('c.png', 3)]
    parts = [part for message in build_input_payload(seen[-1]) for part in message['content']]
    label = next(index for index, part in enumerate(parts) if part.get('text') == '[Ảnh 3: c.png]')
    assert parts[label + 1]['type'] == 'input_image'


async def test_render_queue_waits_in_a_short_line():
    queue = RenderQueue(max_waiting=1)
    order = []

    async def later():
        async with queue.slot(2):
            order.append('sau')

    async with queue.slot(1):
        task = asyncio.create_task(later())
        await asyncio.sleep(.05)
        with pytest.raises(RenderBusy):  # đã có một việc đang chờ: hàng đầy
            async with queue.slot(2): pass
        order.append('trước')
    await task
    assert order == ['trước', 'sau']
    async with queue.slot(1):
        with pytest.raises(RenderBusy):  # chờ quá lâu
            async with queue.slot(.05): pass
    async with queue.slot(.05):
        pass
