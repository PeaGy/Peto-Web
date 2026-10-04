"""Bài thuyết trình PowerPoint (create_presentation): ba phong cách, bảy khuôn, ghi chú, PPTX thật và bản PDF cùng bố cục."""
import json
import re
from io import BytesIO
from types import SimpleNamespace

import pytest
from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE
from pypdf import PdfReader
from PIL import Image

import storage as db
from ai.base import ChatMessage, StreamChunk
from ai.mock import slide_sample
from features.chat import conversations as conversation_actions
from features.chat import service as chat_service
from features.documents.images import prepare
from features.documents.jobs import build_presentation
from features.documents.slides import layout
from features.documents.slides.scene import H, W, TextBox
from features.documents.slides.spec import SCHEMA, DeckInput, load, normalize
from features.documents.tools import DocumentSession, current_session
from conftest import TEST_OWNER, read_events
from test_document_features import attachment, chat, png

EMPTY = dict.fromkeys(('subtitle', 'meta', 'section', 'bullets', 'left_title', 'left_bullets', 'right_title', 'right_bullets',
                       'image', 'caption', 'table_columns', 'table_rows', 'chart_type', 'chart_categories', 'chart_series',
                       'chart_unit'))


def slide(layout_name, title, notes='', **fields):
    return {**EMPTY, 'layout': layout_name, 'title': title, 'notes': notes, **fields}


def full_deck(theme):
    """Đủ bảy khuôn và bốn kiểu biểu đồ."""
    return {'title': 'Hệ thống quản lý thư viện số', 'theme': theme, 'slides': [
        slide('cover', 'Hệ thống quản lý thư viện số', 'Chào hỏi.', subtitle='Đồ án môn Công nghệ phần mềm',
              meta='Nhóm 5 · Lớp KTPM2024\nKhoa Công nghệ thông tin'),
        slide('agenda', 'Nội dung', 'Mục lục.', bullets=['Bài toán', 'Thiết kế', 'Kết quả', 'Kết luận']),
        slide('bullets', 'Bài toán đặt ra', 'Nói về 5–7 phút.', section='1 · Bài toán',
              bullets=['Mượn trả ghi sổ tay, mỗi lượt mất 5–7 phút', 'Không biết sách còn hay đã mượn']),
        slide('two_columns', 'Trước và sau', 'So sánh.', left_title='Hiện tại', left_bullets=['Ghi sổ tay'],
              right_title='Hệ thống mới', right_bullets=['Quét mã vạch', 'Tra cứu trực tuyến']),
        slide('image_text', 'Màn hình tra cứu', 'Mở demo.', image=1, caption='Bản thử nghiệm',
              bullets=['Gõ tên sách hoặc ISBN', 'Đặt giữ sách ngay']),
        slide('table', 'Kết quả thử nghiệm', 'Giải thích bảng.', caption='Đo trên 120 lượt',
              table_columns=['Thao tác', 'Trước', 'Sau', 'Giảm'],
              table_rows=[['Mượn sách', '6,5 phút', '1,2 phút', '82%'], ['Tra cứu', '3,0 phút', '10 giây', '94%']]),
        slide('chart', 'Lượt mượn theo tháng', 'Tháng 9 tăng.', chart_type='column', chart_unit='lượt',
              chart_categories=['T7', 'T8', 'T9'], chart_series=[{'name': '2025', 'values': [190, 240, 410]},
                                                                 {'name': '2026', 'values': [210, 260, 640]}]),
        slide('chart', 'Nguồn độc giả', 'Đa số là sinh viên.', chart_type='pie', chart_categories=['Sinh viên', 'Giảng viên'],
              chart_series=[{'name': 'Độc giả', 'values': [80, 20]}]),
        slide('chart', 'Thời gian thao tác', 'Ngắn hơn.', chart_type='bar', chart_categories=['Mượn', 'Trả'],
              chart_series=[{'name': 'Phút', 'values': [1.2, 0.8]}]),
        slide('chart', 'Đầu sách số hóa', 'Vượt mục tiêu.', chart_type='line', chart_categories=['T7', 'T8', 'T9'],
              chart_series=[{'name': 'Thực tế', 'values': [410, 520, 760]}]),
    ]}


def deck_json(data):
    return normalize(DeckInput.model_validate_json(json.dumps(data, ensure_ascii=False))).to_json()


@pytest.mark.parametrize('theme', ['clean', 'academic', 'bold'])
def test_every_layout_renders_as_real_pptx_and_matching_pdf(theme):
    content = deck_json(full_deck(theme))
    files = build_presentation(content, {1: png()})
    assert files['format'] == 'pptx' and files['pages'] == 10 and files['docx'] == b''
    presentation = Presentation(BytesIO(files['pptx']))
    assert (presentation.slide_width, presentation.slide_height) == (12192000, 6858000)  # 16:9
    slides = list(presentation.slides)
    assert [s.shapes.title.text_frame.text for s in slides][:3] == ['Hệ thống quản lý thư viện số', 'Nội dung', 'Bài toán đặt ra']
    assert all(s.notes_slide.notes_text_frame.text for s in slides)
    assert sum(1 for shape in slides[4].shapes if shape.shape_type == 13) == 1  # ảnh của khuôn ảnh kèm chữ
    table = next(shape.table for shape in slides[5].shapes if shape.has_table)
    assert table.cell(1, 0).text == 'Mượn sách' and table.cell(0, 3).text in ('Giảm', 'GIẢM')
    charts = [next(shape.chart for shape in s.shapes if shape.has_chart) for s in slides[6:]]
    assert [c.chart_type for c in charts] == [XL_CHART_TYPE.COLUMN_CLUSTERED, XL_CHART_TYPE.PIE,
                                              XL_CHART_TYPE.BAR_CLUSTERED, XL_CHART_TYPE.LINE_MARKERS]
    assert list(charts[0].plots[0].series[1].values) == [210, 260, 640]
    assert charts[0].value_axis.maximum_scale >= 640
    assert len(PdfReader(BytesIO(files['pdf'])).pages) == 10
    assert Image.open(BytesIO(files['preview'])).width > 900
    # Mọi ô chữ nằm trong slide: không chữ nào bị đặt ra ngoài khung.
    images = {1: prepare(png())}
    for scene in layout.build(load(content), images):
        for item in scene.items:
            if isinstance(item, TextBox):
                assert 0 <= item.x and item.x + item.w <= W + 0.5 and 0 <= item.y and item.y + item.h <= H + 0.5, item


def test_text_uses_real_bullets_slide_number_fields_and_vietnamese_language():
    files = build_presentation(deck_json(full_deck('clean')), {1: png()})
    presentation = Presentation(BytesIO(files['pptx']))
    xml = presentation.slides[2].shapes._spTree.xml
    assert 'buChar' in xml and 'type="slidenum"' in xml and 'lang="vi-VN"' in xml
    theme_part = presentation.slide_master.part.part_related_by(
        'http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme')
    assert b'0F766E' in theme_part.blob  # màu nhấn của phong cách vào theme tệp


@pytest.mark.parametrize('change, message', [
    (lambda d: d['slides'][2].update(bullets=['ý'] * 7), 'Slide 3 (bullets): bullets cần 1–6 ý'),
    (lambda d: d['slides'][5].update(table_rows=[['a', 'b']]), 'Slide 6 (table): mỗi hàng phải có đúng 4 ô'),
    (lambda d: d['slides'][7]['chart_series'].append({'name': 'Thêm', 'values': [1, 2]}), 'Slide 8 (chart): biểu đồ tròn chỉ có một chuỗi'),
    (lambda d: d['slides'][4].update(image=None), 'Slide 5 (image_text): image phải là số N'),
    (lambda d: d['slides'][6]['chart_series'][0].update(values=[1, 2]), 'phải có đúng 3 số'),
])
def test_validation_names_the_slide_and_what_to_fix(change, message):
    data = full_deck('clean')
    change(data)
    with pytest.raises(ValueError, match=re.escape(message)):
        normalize(DeckInput.model_validate_json(json.dumps(data, ensure_ascii=False)))


def test_content_that_cannot_fit_asks_to_split_the_slide():
    data = full_deck('bold')
    cell = 'Một ô bảng dài gần hết giới hạn sáu mươi ký tự cho phép nhé'
    data['slides'][5].update(table_columns=['A', 'B', 'C', 'D', 'E'], table_rows=[[cell] * 5 for _ in range(7)])
    with pytest.raises(layout.SlideOverflow, match=re.escape('Slide 6 (table): Bảng quá dài')):
        build_presentation(deck_json(data), {1: png()})


def test_schema_is_strict_and_lists_every_field():
    item = SCHEMA['parameters']['properties']['slides']['items']
    assert SCHEMA['strict'] and set(item['required']) == set(item['properties'])
    assert item['properties']['chart_type']['enum'][-1] is None
    assert 'không bịa số liệu' in SCHEMA['description']


async def test_slide_request_creates_pptx_card_downloads_and_notes(client, anon_client, monkeypatch):
    events = await read_events(await client.post('/api/chat', json={'message': '__slide__:academic'}))
    assert events[-1]['type'] == 'done'
    artifact = next(event['artifact'] for event in events if event['type'] == 'artifact')
    conversation = next(event['conversation_id'] for event in events if event['type'] == 'meta')
    assert artifact['format'] == 'pptx' and artifact['style'] == 'academic' and artifact['pages'] == 6
    assert artifact['filename'].endswith('.pptx')
    base = f"/api/documents/{artifact['id']}"
    listing = (await client.get(f'/api/documents?conversation_id={conversation}')).json()['documents']
    assert listing[0]['format'] == 'pptx'
    detail = (await client.get(base + '?version=1')).json()
    assert json.loads(detail['content'])['slides'][1]['notes'] == 'Đi nhanh qua mục lục.'
    deck = await client.get(base + '/export/pptx?version=1')
    assert deck.headers['content-type'] == 'application/vnd.openxmlformats-officedocument.presentationml.presentation'
    assert len(Presentation(BytesIO(deck.content)).slides) == 6
    pdf = await client.get(base + '/export/pdf?version=1')
    assert len(PdfReader(BytesIO(pdf.content)).pages) == 6
    assert (await client.get(base + '/export/docx?version=1')).status_code == 400
    assert (await client.get(base + '/preview?version=1&page=6')).status_code == 200
    assert (await anon_client.get(base + '/export/pptx?version=1')).status_code == 401
    edit = await client.post(base + '/versions', json={'title': 'Sửa', 'content': '# Sửa', 'base_version': 1})
    assert edit.status_code == 400 and 'Nhờ Peto sửa' in edit.json()['detail']
    seen = []

    class FollowUp:
        async def stream(self, **kwargs):
            seen.extend(kwargs['messages'])
            yield 'Đã đọc bài cũ.'
    monkeypatch.setattr(chat_service, 'get_provider', lambda model="peto": FollowUp())
    await client.post('/api/chat', json={'message': 'Sửa slide 3', 'conversation_id': conversation})
    # Peto nhận lại bài cũ (JSON gọn) để sửa ở lượt sau.
    assert any('"layout": "two_columns"' in a.text_excerpt for m in seen for a in m.attachments)


async def test_word_documents_have_no_pptx_download(client):
    events = await read_events(await client.post('/api/chat', json={'message': 'Tạo file docx viết một bài nghị luận xã hội'}))
    artifact = next(event['artifact'] for event in events if event['type'] == 'artifact')
    response = await client.get(f"/api/documents/{artifact['id']}/export/pptx?version=1")
    assert response.status_code == 400


async def test_images_come_only_from_the_conversation(client):
    conversation = await chat(client, 'Ảnh màn hình', images=[attachment('man-hinh.png', png())])
    session = DocumentSession(TEST_OWNER, conversation)
    data = full_deck('clean')
    result = await session.present(json.dumps(data, ensure_ascii=False))
    assert result['ok'], result
    data['slides'][4]['image'] = 2
    missing = await session.present(json.dumps(data, ensure_ascii=False))
    assert 'không có Ảnh 2' in missing['error'] and 'create_presentation' in missing['error']
    other = DocumentSession('someone-else', await db.create_conversation('someone-else', 'Khác'))
    assert 'chưa có ảnh' in (await other.present(json.dumps(full_deck('clean'), ensure_ascii=False)))['error']


async def test_unaccented_vietnamese_slides_are_rejected(client):
    conversation = await db.create_conversation(TEST_OWNER, 'Không dấu')
    session = DocumentSession(TEST_OWNER, conversation)
    data = slide_sample('clean')
    data['title'] = 'Nghi luan xa hoi'
    data['slides'] = [slide('bullets', 'Trach nhiem cua gioi tre', bullets=['Doc sach trong thoi dai so', 'Nguyen nhan va giai phap'])]
    result = await session.present(json.dumps(data))
    assert 'không dấu' in result['error'] and not session.created


async def test_branch_keeps_the_presentation_file(client):
    conversation = await db.create_conversation(TEST_OWNER, 'Nhánh slide')
    await db.add_message(conversation, 'user', 'Làm slide')
    session = DocumentSession(TEST_OWNER, conversation)
    result = await session.present(json.dumps(slide_sample('bold'), ensure_ascii=False))
    await db.add_message(conversation, 'assistant', 'Đã tạo', artifacts=[result['artifact']])
    second = await db.add_message(conversation, 'user', 'Sửa lại')
    branch, _ = await conversation_actions.fork(TEST_OWNER, conversation, second, 'Sửa lại giúp mình')
    copied = (await db.get_messages(TEST_OWNER, branch))[1]['artifacts'][0]
    assert copied['id'] != result['artifact']['id']
    response = await client.get(f"/api/documents/{copied['id']}/export/pptx?version=1")
    assert response.status_code == 200 and len(Presentation(BytesIO(response.content)).slides) == 6


async def test_real_provider_offers_the_presentation_tool(client, monkeypatch):
    from test_clock_tools import FakeStream, call, done, fake_provider
    arguments = json.dumps(slide_sample('clean'), ensure_ascii=False)
    first = FakeStream([done(call('create_presentation', arguments))])
    second = FakeStream([SimpleNamespace(type='response.output_text.delta', delta='Đã tạo slide.'), done()])
    provider, requests = fake_provider(monkeypatch, [first, second])
    conversation = await db.create_conversation(TEST_OWNER, 'Slide qua provider')
    token = current_session.set(DocumentSession(TEST_OWNER, conversation))
    try:
        chunks = [chunk async for chunk in provider.stream(system_prompt='Peto', messages=[ChatMessage('user', 'Làm slide')])]
    finally:
        current_session.reset(token)
    assert any(tool.get('name') == 'create_presentation' for tool in requests[0]['tools'])
    result = json.loads(requests[1]['input'][-1]['output'])
    assert result['ok'] is True and result['artifact']['format'] == 'pptx'
    assert any(isinstance(chunk, StreamChunk) and chunk.kind == 'document_status' and 'slide' in chunk.text for chunk in chunks)
