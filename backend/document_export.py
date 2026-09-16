"""Render bounded Markdown to DOCX/PDF without executing HTML or fetching assets."""
from __future__ import annotations

from dataclasses import dataclass, field
from html import escape
from io import BytesIO
from pathlib import Path
import re
import threading
import unicodedata

from markdown_it import MarkdownIt

MAX_CONTENT = 60000
FONT_DIR = Path(__file__).parent / "assets" / "fonts"
_font_lock = threading.Lock()


@dataclass
class Run:
    text: str
    bold: bool = False
    italic: bool = False
    strike: bool = False
    code: bool = False


@dataclass
class Block:
    kind: str
    runs: list[Run] = field(default_factory=list)
    level: int = 0
    prefix: str = ""
    rows: list[list[list[Run]]] = field(default_factory=list)


def clean_text(value: str) -> str:
    value = unicodedata.normalize("NFC", value.replace("\r\n", "\n").replace("\r", "\n"))
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f\ud800-\udfff\ufffe\uffff]", "", value).strip()


def inline_runs(tokens) -> list[Run]:
    runs, bold, italic, strike = [], 0, 0, 0
    links = []
    for token in tokens or []:
        if token.type == "strong_open": bold += 1
        elif token.type == "strong_close": bold -= 1
        elif token.type == "em_open": italic += 1
        elif token.type == "em_close": italic -= 1
        elif token.type == "s_open": strike += 1
        elif token.type == "s_close": strike -= 1
        elif token.type in {"text", "code_inline", "softbreak", "hardbreak", "image"}:
            text = token.content
            if token.type in {"softbreak", "hardbreak"}: text = "\n"
            if token.type == "image": text = f"[Ảnh: {token.content or 'không có mô tả'}]"
            runs.append(Run(text, bool(bold), bool(italic), bool(strike), token.type == "code_inline"))
        elif token.type == "link_close":
            if links:
                url, start = links.pop()
                label = ''.join(run.text for run in runs[start:])
                if url and url != label:
                    runs.append(Run(f' ({url})'))
        elif token.type == "link_open":
            # Keep the source address in downloads without fetching or executing it.
            links.append((token.attrGet('href'), len(runs)))
    return runs


def parse_blocks(content: str) -> list[Block]:
    if not content.strip() or len(content) > MAX_CONTENT:
        raise ValueError("Nội dung cần có từ 1 đến 60.000 ký tự.")
    tokens = MarkdownIt("commonmark", {"html": False, "maxNesting": 20}).enable(["table", "strikethrough"]).parse(content)
    if len(tokens) > 12000:
        raise ValueError("Tài liệu có quá nhiều mục. Hãy chia thành các tài liệu nhỏ hơn.")
    blocks, lists, pending, index, quote_depth = [], [], "", 0, 0
    while index < len(tokens):
        token = tokens[index]
        if token.type in {"bullet_list_open", "ordered_list_open"}:
            lists.append([token.type == "ordered_list_open", int(token.attrGet("start") or 1)])
        elif token.type in {"bullet_list_close", "ordered_list_close"}:
            lists.pop(); pending = ""
        elif token.type == "list_item_open":
            ordered, number = lists[-1]
            pending = f"{number}. " if ordered else "• "
            lists[-1][1] += 1
        elif token.type == "list_item_close": pending = ""
        elif token.type == "blockquote_open": quote_depth += 1
        elif token.type == "blockquote_close": quote_depth -= 1
        elif token.type in {"paragraph_open", "heading_open"}:
            inline = tokens[index + 1]
            kind = "heading" if token.type == "heading_open" else "quote" if quote_depth else "paragraph"
            level = int(token.tag[1:]) if kind == "heading" else min(len(lists), 6)
            blocks.append(Block(kind, inline_runs(inline.children), level, pending))
            pending = ""
        elif token.type in {"fence", "code_block"}:
            for line in token.content.rstrip("\n").split("\n"):
                blocks.append(Block("code", [Run(line or " ", code=True)]))
        elif token.type == "hr": blocks.append(Block("rule"))
        elif token.type == "table_open":
            rows, row = [], []
            index += 1
            while index < len(tokens) and tokens[index].type != "table_close":
                cell = tokens[index]
                if cell.type == "tr_open": row = []
                elif cell.type == "inline": row.append(inline_runs(cell.children))
                elif cell.type == "tr_close": rows.append(row)
                index += 1
            if len(rows) > 100 or any(len(row) > 8 for row in rows):
                raise ValueError("Mỗi bảng hỗ trợ tối đa 100 dòng và 8 cột. Hãy chia bảng lớn trước khi xuất.")
            blocks.append(Block("table", rows=rows))
        index += 1
    if len(blocks) > 1200:
        raise ValueError("Tài liệu có quá nhiều đoạn. Hãy chia nhỏ trước khi xuất.")
    return blocks


def _body(title, content):
    blocks = parse_blocks(content)
    if blocks and blocks[0].kind == "heading" and "".join(run.text for run in blocks[0].runs).strip() == title.strip():
        blocks.pop(0)
    return blocks


def render_docx(title: str, content: str, layout: str = 'report') -> bytes:
    from docx import Document
    from docx.shared import Inches, Cm, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    document = Document()
    section = document.sections[0]
    essay = layout == 'essay'
    section.page_width, section.page_height = (Cm(21), Cm(29.7)) if essay else (Inches(8.5), Inches(11))
    section.top_margin = section.bottom_margin = Inches(.75)
    section.left_margin = section.right_margin = Inches(.75)
    for name in ['Normal', 'Title', *[f'Heading {n}' for n in range(1, 7)]]:
        style = document.styles[name]
        style.font.name = 'Times New Roman' if essay else 'Arial'
        style.font.color.rgb = RGBColor(0, 0, 0)
    normal = document.styles['Normal']
    normal.font.size = Pt(13 if essay else 11)
    normal.paragraph_format.line_spacing = 1.5 if essay else 1.25
    if essay:
        normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        normal.paragraph_format.first_line_indent = Cm(.75)
        header = section.header.paragraphs[0]
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        header.add_run('Nghị luận xã hội').italic = True
    normal.paragraph_format.space_after = Pt(7)
    document.styles['Title'].font.size = Pt(24)
    document.core_properties.title, document.core_properties.author = title, 'Peto'
    title_paragraph = document.add_paragraph(title, 'Title')
    title_paragraph.paragraph_format.first_line_indent = 0
    if essay: title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def write(paragraph, runs, prefix=''):
        if prefix: paragraph.add_run(prefix)
        for item in runs:
            run = paragraph.add_run(item.text)
            run.bold = item.bold or None
            run.italic, run.font.strike = item.italic or None, item.strike or None
            if item.code: run.font.name, run.font.size = 'Consolas', Pt(10)

    for block in _body(title, content):
        if block.kind == 'table' and block.rows:
            table = document.add_table(rows=0, cols=len(block.rows[0]))
            table.style = 'Table Grid'
            for row_index, row in enumerate(block.rows):
                cells = table.add_row().cells
                for cell, runs in zip(cells, row):
                    write(cell.paragraphs[0], runs)
                    if row_index == 0:
                        for run in cell.paragraphs[0].runs: run.bold = True
                        shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'EDF0F4'); cell._tc.get_or_add_tcPr().append(shade)
                if row_index == 0:
                    repeat = OxmlElement('w:tblHeader'); table.rows[0]._tr.get_or_add_trPr().append(repeat)
            document.add_paragraph().paragraph_format.space_after = Pt(3)
        elif block.kind == 'rule': document.add_paragraph()
        else:
            style = f'Heading {block.level}' if block.kind == 'heading' else 'Normal'
            paragraph = document.add_paragraph(style=style)
            if block.kind == 'heading':
                paragraph.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
                paragraph.paragraph_format.first_line_indent = 0
            if block.level and block.kind != 'heading': paragraph.paragraph_format.left_indent = Inches(.2 * block.level)
            if block.kind == 'quote': paragraph.paragraph_format.left_indent = Inches(.25)
            if block.kind == 'code': paragraph.paragraph_format.space_after = Pt(0)
            write(paragraph, block.runs, block.prefix)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('Trang ').font.size = Pt(9)
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
    output = BytesIO(); document.save(output)
    return output.getvalue()


def _register_fonts():
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    with _font_lock:
        if 'PetoSans' not in pdfmetrics.getRegisteredFontNames():
            for suffix, style in [('', 'Regular'), ('-Bold', 'Bold'), ('-Italic', 'Italic'), ('-BoldItalic', 'BoldItalic')]:
                pdfmetrics.registerFont(TTFont('PetoSans' + suffix, str(FONT_DIR / f'NotoSans-{style}.ttf')))
            pdfmetrics.registerFontFamily('PetoSans', normal='PetoSans', bold='PetoSans-Bold', italic='PetoSans-Italic', boldItalic='PetoSans-BoldItalic')
        if 'PetoSerif' not in pdfmetrics.getRegisteredFontNames():
            for suffix, style in [('', 'Regular'), ('-Bold', 'Bold'), ('-Italic', 'Italic'), ('-BoldItalic', 'BoldItalic')]:
                pdfmetrics.registerFont(TTFont('PetoSerif' + suffix, str(FONT_DIR / f'NotoSerif-{style}.ttf')))
            pdfmetrics.registerFontFamily('PetoSerif', normal='PetoSerif', bold='PetoSerif-Bold', italic='PetoSerif-Italic', boldItalic='PetoSerif-BoldItalic')


def render_pdf(title: str, content: str, layout: str = 'report') -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter, A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, LongTable, TableStyle, HRFlowable
    from reportlab.pdfbase import pdfmetrics

    _register_fonts()
    essay = layout == 'essay'
    font = 'PetoSerif' if essay else 'PetoSans'
    page_size = A4 if essay else letter
    if any(not char.isspace() and ord(char) not in pdfmetrics.getFont(font).face.charToGlyph for char in title + content):
        raise ValueError('PDF chưa hỗ trợ một số ký tự trong bản nháp (chẳng hạn emoji hoặc chữ tượng hình). Hãy bỏ các ký tự đó hoặc tải DOCX.')
    body = ParagraphStyle('Body', fontName='PetoSans', fontSize=11, leading=16, spaceAfter=8, splitLongWords=True)
    heading = {n: ParagraphStyle(f'Heading{n}', parent=body, fontName='PetoSans-Bold', fontSize=max(11, 20-2*n), leading=max(16, 25-2*n), spaceBefore=12, keepWithNext=True) for n in range(1, 7)}
    title_style = ParagraphStyle('Title', parent=body, fontName='PetoSans-Bold', fontSize=24, leading=31, spaceAfter=18)
    if essay:
        body.fontName, body.fontSize, body.leading = font, 13, 20
        body.alignment, body.firstLineIndent = 4, 21
        for style in heading.values(): style.fontName = font + '-Bold'
        title_style.fontName, title_style.fontSize, title_style.leading = font + '-Bold', 21, 29
        title_style.alignment, title_style.spaceAfter = 1, 26

    def markup(runs):
        chunks = []
        for run in runs:
            text = escape(run.text).replace('\n', '<br/>')
            if run.code: text = text.replace(' ', '&#160;')
            if run.bold: text = f'<b>{text}</b>'
            if run.italic: text = f'<i>{text}</i>'
            if run.strike: text = f'<strike>{text}</strike>'
            chunks.append(text)
        return ''.join(chunks) or '&#160;'

    story = [Paragraph(escape(title), title_style)]
    for block in _body(title, content):
        if block.kind == 'table' and block.rows:
            cell_style = ParagraphStyle('Cell', parent=body, fontSize=10, leading=14, spaceAfter=0)
            head_style = ParagraphStyle('CellHead', parent=cell_style, fontName=font + '-Bold')
            rows = [[Paragraph(markup(cell), head_style if i == 0 else cell_style) for cell in row] for i, row in enumerate(block.rows)]
            table = LongTable(rows, colWidths=[(page_size[0] - 108) / len(rows[0])] * len(rows[0]), repeatRows=1, splitByRow=1, splitInRow=1)
            table.setStyle(TableStyle([
                ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#EDF0F4')),
                ('GRID', (0,0), (-1,-1), .5, colors.HexColor('#BDC5CE')),
                ('VALIGN', (0,0), (-1,-1), 'TOP'),
                ('LEFTPADDING', (0,0), (-1,-1), 8), ('RIGHTPADDING', (0,0), (-1,-1), 8),
                ('TOPPADDING', (0,0), (-1,-1), 7), ('BOTTOMPADDING', (0,0), (-1,-1), 7),
            ]))
            story.extend([table, Spacer(1, 10)])
        elif block.kind == 'rule': story.append(HRFlowable(width='100%', thickness=.5, color=colors.HexColor('#BDC5CE'), spaceAfter=10))
        else:
            style = heading[block.level] if block.kind == 'heading' else ParagraphStyle('Item', parent=body, leftIndent=14 * block.level)
            if block.kind == 'quote': style.leftIndent = 18
            if block.kind == 'code': style.fontSize, style.leading, style.spaceAfter = 10, 14, 0
            story.append(Paragraph(escape(block.prefix) + markup(block.runs), style))
    def page_footer(canvas, document):
        canvas.saveState(); canvas.setFont(font, 9)
        canvas.drawCentredString(page_size[0] / 2, 30, f'Trang {document.page}')
        if essay:
            canvas.setFont(font + '-Italic', 9)
            canvas.setFillColor(colors.HexColor('#506279'))
            canvas.drawRightString(page_size[0] - 54, page_size[1] - 34, 'Nghị luận xã hội')
        canvas.restoreState()
    output = BytesIO()
    document = SimpleDocTemplate(output, pagesize=page_size, topMargin=62 if essay else 54, bottomMargin=54, leftMargin=54, rightMargin=54, title=title, author='Peto')
    document.build(story, onFirstPage=page_footer, onLaterPages=page_footer)
    return output.getvalue()
