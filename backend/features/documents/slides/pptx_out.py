"""Vẽ các Scene thành tệp PPTX thật bằng python-pptx.

Tiêu đề nằm trong ô tiêu đề thật của slide (hiện trong outline, đọc được bằng trình đọc màn hình), gạch đầu dòng là gạch đầu
dòng của PowerPoint (gõ thêm dòng vẫn có dấu), bảng và biểu đồ là đối tượng gốc sửa được số liệu, số trang là trường tự đổi
khi sắp xếp lại slide. Mọi ô chữ đặt lề trong bằng 0 và khoảng dòng cố định, để chữ nằm đúng chỗ như bản xem trước.
"""
from __future__ import annotations

from io import BytesIO
import uuid

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_MARKER_STYLE, XL_TICK_MARK
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from features.documents.slides.fonts import SANS
from features.documents.slides.scene import H, W, Chart, Line, Para, Picture, Scene, Shape, Table, TextBox
from features.documents.slides.themes import Theme

POINT = 12700  # EMU mỗi point
NO_STYLE_TABLE = '{2D5ABB26-0587-4C30-8999-92F81FD0307C}'  # kiểu bảng "No Style, No Grid" của Office
_ALIGN = {'left': PP_ALIGN.LEFT, 'right': PP_ALIGN.RIGHT, 'center': PP_ALIGN.CENTER}
_ANCHOR = {'top': MSO_ANCHOR.TOP, 'middle': MSO_ANCHOR.MIDDLE, 'bottom': MSO_ANCHOR.BOTTOM}
_A = 'http://schemas.openxmlformats.org/drawingml/2006/main'


def emu(value: float) -> Emu:
    return Emu(int(round(value * POINT)))


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def _plain(shape):
    """Bỏ kiểu mặc định của theme (bóng, viền theo theme), màu đã đặt rõ ở từng hình."""
    style = shape._element.find(qn('p:style'))
    if style is not None:
        shape._element.remove(style)


def _raise_to_top(shape):
    """Đưa hình lên trên cùng (ô tiêu đề có sẵn từ đầu, phải nằm trên dải màu vẽ sau nó)."""
    element = shape._element
    tree = element.getparent()
    tree.remove(element)
    extension = tree.find(qn('p:extLst'))
    if extension is not None:
        extension.addprevious(element)
    else:
        tree.append(element)


def _font(run, item: Para):
    font = run.font
    font.name, font.size, font.bold, font.italic = item.family, Pt(item.size), item.bold, item.italic
    font.color.rgb = rgb(item.color)
    properties = run._r.get_or_add_rPr()
    properties.set('lang', 'vi-VN')
    if item.spacing:
        properties.set('spc', str(int(round(item.spacing * 100))))


def _bullet(paragraph, item: Para):
    properties = paragraph._p.get_or_add_pPr()
    for tag in ('a:buClr', 'a:buSzPct', 'a:buFont', 'a:buChar', 'a:buNone', 'a:buAutoNum'):
        for old in properties.findall(qn(tag)):
            properties.remove(old)
    if not item.bullet:
        etree.SubElement(properties, qn('a:buNone'))
        return
    mark = item.bullet
    properties.set('marL', str(int(mark.indent * POINT)))
    properties.set('indent', str(int(-mark.indent * POINT)))
    color = etree.SubElement(properties, qn('a:buClr'))
    etree.SubElement(color, qn('a:srgbClr')).set('val', mark.color)
    etree.SubElement(properties, qn('a:buSzPct')).set('val', str(int(mark.scale * 100000)))
    etree.SubElement(properties, qn('a:buFont')).set('typeface', SANS)
    etree.SubElement(properties, qn('a:buChar')).set('char', mark.char)


def _frame(frame, box_anchor: str = 'top'):
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = _ANCHOR[box_anchor]


def _paragraphs(frame, paras: list[Para], slide_number: bool = False):
    for index, item in enumerate(paras):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = _ALIGN[item.align]
        paragraph.line_spacing = Pt(item.line)
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(item.after)
        _bullet(paragraph, item)
        run = paragraph.add_run()
        run.text = item.text
        _font(run, item)
        if slide_number:
            # Trường số slide: PowerPoint tự ghi số đúng khi thêm, bớt hay đổi thứ tự slide.
            field = run._r
            field.tag = qn('a:fld')
            field.set('id', '{' + str(uuid.uuid4()).upper() + '}')
            field.set('type', 'slidenum')


def _text(slide, box: TextBox, placeholder=None):
    shape = placeholder or slide.shapes.add_textbox(emu(box.x), emu(box.y), emu(box.w), emu(max(box.h, 1)))
    if placeholder is not None:
        shape.left, shape.top, shape.width, shape.height = emu(box.x), emu(box.y), emu(box.w), emu(max(box.h, 1))
        shape.text_frame.text = ''
        _raise_to_top(shape)
    _frame(shape.text_frame, box.anchor)
    _paragraphs(shape.text_frame, box.paras, box.slide_number)


def _shape(slide, item: Shape):
    kind = {'rect': MSO_SHAPE.RECTANGLE, 'round': MSO_SHAPE.ROUNDED_RECTANGLE, 'ellipse': MSO_SHAPE.OVAL}[item.kind]
    shape = slide.shapes.add_shape(kind, emu(item.x), emu(item.y), emu(item.w), emu(item.h))
    _plain(shape)
    if item.kind == 'round':
        shape.adjustments[0] = min(0.5, item.radius / min(item.w, item.h))
    if item.fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(item.fill)
    else:
        shape.fill.background()
    if item.line:
        shape.line.color.rgb = rgb(item.line)
        shape.line.width = Pt(item.line_width)
    else:
        shape.line.fill.background()


def _line(slide, item: Line):
    connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, emu(item.x1), emu(item.y1), emu(item.x2), emu(item.y2))
    _plain(connector)
    connector.line.color.rgb = rgb(item.color)
    connector.line.width = Pt(item.width)
    if item.dash:
        connector.line.dash_style = MSO_LINE_DASH_STYLE.ROUND_DOT


def _picture(slide, item: Picture):
    picture = slide.shapes.add_picture(BytesIO(item.data), emu(item.x), emu(item.y), emu(item.w), emu(item.h))
    if item.border:
        picture.line.color.rgb = rgb(item.border)
        picture.line.width = Pt(0.75)


def _border(cell, bottom: tuple[str, float] | None):
    properties = cell._tc.get_or_add_tcPr()
    for tag in ('a:lnL', 'a:lnR', 'a:lnT', 'a:lnB'):
        for old in properties.findall(qn(tag)):
            properties.remove(old)
    for index, (tag, spec) in enumerate((('a:lnL', None), ('a:lnR', None), ('a:lnT', None), ('a:lnB', bottom))):
        line = etree.Element(qn(tag))
        if spec:
            line.set('w', str(int(spec[1] * POINT)))
            fill = etree.SubElement(line, qn('a:solidFill'))
            etree.SubElement(fill, qn('a:srgbClr')).set('val', spec[0])
        else:
            line.set('w', '0')
            etree.SubElement(line, qn('a:noFill'))
        properties.insert(index, line)


def _table(slide, item: Table):
    rows, columns = len(item.cells), len(item.cells[0])
    frame = slide.shapes.add_table(rows, columns, emu(item.x), emu(item.y), emu(sum(item.widths)), emu(sum(item.heights)))
    table = frame.table
    table.first_row = True
    table.horz_banding = False
    style = frame._element.graphic.graphicData.tbl.tblPr.find(qn('a:tableStyleId'))
    if style is not None:
        style.text = NO_STYLE_TABLE
    for index, width in enumerate(item.widths):
        table.columns[index].width = emu(width)
    for index, height in enumerate(item.heights):
        table.rows[index].height = emu(height)
    for r, row in enumerate(item.cells):
        for c, cell in enumerate(row):
            target = table.cell(r, c)
            target.margin_left = target.margin_right = emu(item.pad_x)
            target.margin_top = target.margin_bottom = emu(item.pad_y)
            target.vertical_anchor = MSO_ANCHOR.MIDDLE
            if cell.fill:
                target.fill.solid()
                target.fill.fore_color.rgb = rgb(cell.fill)
            else:
                target.fill.background()
            frame_text = target.text_frame
            frame_text.word_wrap = True
            _paragraphs(frame_text, [cell.para])
            _border(target, item.header_rule if r == 0 and item.header_rule else item.row_rule)


def _chart(slide, item: Chart):
    kind = {'column': XL_CHART_TYPE.COLUMN_CLUSTERED, 'bar': XL_CHART_TYPE.BAR_CLUSTERED,
            'line': XL_CHART_TYPE.LINE_MARKERS, 'pie': XL_CHART_TYPE.PIE}[item.kind]
    categories, series = item.categories, item.series
    if item.kind == 'bar':
        # PowerPoint vẽ nhãn đầu tiên ở dưới cùng; đảo thứ tự để nhãn đầu nằm trên cùng như bản xem trước.
        categories = categories[::-1]
        series = [(name, values[::-1]) for name, values in series]
    data = CategoryChartData(number_format='#,##0.##')
    data.categories = categories
    for name, values in series:
        data.add_series(name, values)
    chart = slide.shapes.add_chart(kind, emu(item.x), emu(item.y), emu(item.w), emu(item.h), data).chart
    chart.has_title = False
    chart.font.name, chart.font.size = item.family, Pt(12)
    chart.font.color.rgb = rgb(item.text)
    plot = chart.plots[0]
    if item.kind == 'pie':
        chart.has_legend = True
        chart.legend.position = XL_LEGEND_POSITION.RIGHT
        chart.legend.include_in_layout = False
        plot.has_data_labels = True
        labels = plot.data_labels
        labels.show_percentage, labels.show_value, labels.show_category_name = True, False, False
        labels.number_format, labels.number_format_is_linked = '0%', False
        labels.position = XL_LABEL_POSITION.CENTER
        labels.font.size, labels.font.bold = Pt(12), True
        labels.font.color.rgb = rgb('FFFFFF')
        for index, point in enumerate(plot.series[0].points):
            point.format.fill.solid()
            point.format.fill.fore_color.rgb = rgb(item.colors[index])
            point.format.line.color.rgb = rgb('FFFFFF')
        return
    chart.has_legend = len(item.series) > 1
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.TOP
        chart.legend.include_in_layout = False
    value = chart.value_axis
    value.minimum_scale, value.maximum_scale, value.major_unit = item.minimum, item.maximum, item.step
    value.has_major_gridlines = True
    value.major_gridlines.format.line.color.rgb = rgb(item.grid)
    value.major_gridlines.format.line.width = Pt(0.75)
    value.format.line.fill.background()
    value.major_tick_mark = XL_TICK_MARK.NONE
    value.tick_labels.number_format, value.tick_labels.number_format_is_linked = '#,##0.##', False
    category = chart.category_axis
    category.format.line.color.rgb = rgb(item.grid)
    category.major_tick_mark = XL_TICK_MARK.NONE
    if item.kind in ('column', 'bar'):
        plot.gap_width = 70
        plot.overlap = -12 if len(item.series) > 1 else 0
    for index, entry in enumerate(plot.series):
        color = rgb(item.colors[index])
        if item.kind == 'line':
            entry.smooth = False
            entry.format.line.color.rgb = color
            entry.format.line.width = Pt(2.5)
            entry.marker.style, entry.marker.size = XL_MARKER_STYLE.CIRCLE, 7
            entry.marker.format.fill.solid()
            entry.marker.format.fill.fore_color.rgb = color
            entry.marker.format.line.color.rgb = color
        else:
            entry.format.fill.solid()
            entry.format.fill.fore_color.rgb = color
            entry.format.line.fill.background()
    if item.labels:
        plot.has_data_labels = True
        labels = plot.data_labels
        labels.number_format, labels.number_format_is_linked = '#,##0.##', False
        labels.font.size = Pt(12)
        labels.font.color.rgb = rgb(item.text)
        labels.position = XL_LABEL_POSITION.OUTSIDE_END


def _theme(presentation, theme: Theme):
    """Phông và màu của theme tệp: chữ người dùng gõ thêm hay biểu đồ chèn thêm trong PowerPoint vẫn hợp phong cách."""
    part = presentation.slide_master.part.part_related_by(RT.THEME)
    root = etree.fromstring(part.blob)
    names = {'a': _A}
    for path, face in (('.//a:majorFont/a:latin', theme.heading), ('.//a:minorFont/a:latin', SANS)):
        node = root.find(path, names)
        if node is not None:
            node.set('typeface', face)
    palette = [theme.accent, *theme.pie[1:4], theme.muted, theme.hair]
    for index, color in enumerate(palette[:6], start=1):
        node = root.find(f'.//a:clrScheme/a:accent{index}', names)
        if node is not None:
            for child in list(node):
                node.remove(child)
            etree.SubElement(node, f'{{{_A}}}srgbClr').set('val', color)
    part._blob = etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)


def _widen(presentation, factor: float):
    """Mẫu mặc định của python-pptx là 4:3; kéo các ô của mẫu theo chiều ngang để slide người dùng thêm vẫn cân đối."""
    owners = [presentation.slide_master, *presentation.slide_layouts]
    for owner in owners:
        for shape in owner.placeholders:
            if shape._element.xpath('./p:spPr/a:xfrm'):
                shape.left, shape.width = int(shape.left * factor), int(shape.width * factor)


def render(scenes: list[Scene], title: str, theme: Theme) -> bytes:
    presentation = Presentation()
    old_width = presentation.slide_width
    presentation.slide_width, presentation.slide_height = emu(W), emu(H)
    _widen(presentation, presentation.slide_width / old_width)
    _theme(presentation, theme)
    properties = presentation.core_properties
    properties.title, properties.last_modified_by, properties.language = title, 'Peto', 'vi-VN'
    layout = presentation.slide_layouts[5]  # Title Only: có ô tiêu đề thật
    draw = {Shape: _shape, Line: _line, Picture: _picture, Table: _table, Chart: _chart}
    for scene in scenes:
        slide = presentation.slides.add_slide(layout)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = rgb(scene.background)
        placeholder = slide.shapes.title
        used = False
        for item in scene.items:
            if isinstance(item, TextBox):
                if item.title and not used and placeholder is not None:
                    _text(slide, item, placeholder)
                    used = True
                else:
                    _text(slide, item)
            else:
                draw[type(item)](slide, item)
        if placeholder is not None and not used:
            placeholder._element.getparent().remove(placeholder._element)
        if scene.notes:
            slide.notes_slide.notes_text_frame.text = scene.notes
    output = BytesIO()
    presentation.save(output)
    return output.getvalue()
