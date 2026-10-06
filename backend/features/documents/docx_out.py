"""Bản Word (DOCX) của tài liệu Peto tạo, theo kiểu trình bày (themes.py), bằng python-docx.

Bìa dựng bằng một bảng ẩn viền ba hàng (trên, giữa, dưới) có chiều cao cố định, vì Word không có "đẩy xuống đáy trang":
khối tên trường nằm trên, tên đề tài giữa trang, thông tin người làm ở đáy. Khung đôi của kiểu Khung đôi là viền trang
chỉ ở trang đầu (w:pgBorders display="firstPage"); dải màu của kiểu Dải màu là một hình chữ nhật neo theo trang.
Các phần tử XML tự thêm được chèn đúng thứ tự lược đồ (_insert), vì Word báo tệp hỏng khi thứ tự sai.
"""
from __future__ import annotations

import re
from io import BytesIO

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Pt, RGBColor

from features.documents import export as base
from features.documents.cover import DOTS, Cover
from features.documents.themes import NAVY, Theme, theme_for

FONT = "Times New Roman"
TEXT_WIDTH = 16.0  # cm: A4 rộng 21 cm, lề trái 3, lề phải 2
TWIPS = 566.929  # twip mỗi cm
GRAY = "5A6B80"
_ORDER = {
    "pPr": ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr",
            "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap", "overflowPunct",
            "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd", "snapToGrid", "spacing", "ind",
            "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc", "textDirection", "textAlignment",
            "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr", "pPrChange"],
    "rPr": ["rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps", "strike", "dstrike", "outline", "shadow",
            "emboss", "imprint", "noProof", "snapToGrid", "vanish", "webHidden", "color", "spacing", "w", "kern",
            "position", "sz", "szCs", "highlight", "u", "effect", "bdr", "shd", "fitText", "vertAlign", "rtl", "cs",
            "em", "lang", "eastAsianLayout", "specVanish", "oMath"],
    "tcPr": ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar",
             "textDirection", "tcFitText", "vAlign", "hideMark"],
    "tblPr": ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize", "tblW",
              "jc", "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook"],
    "trPr": ["cnfStyle", "divId", "gridBefore", "gridAfter", "wBefore", "wAfter", "cantSplit", "trHeight",
             "tblHeader", "tblCellSpacing", "jc", "hidden"],
    "sectPr": ["headerReference", "footerReference", "footnotePr", "endnotePr", "type", "pgSz", "pgMar", "paperSrc",
               "pgBorders", "lnNumType", "pgNumType", "cols", "formProt", "vAlign", "noEndnote", "titlePg",
               "textDirection", "bidi", "rtlGutter", "docGrid", "printerSettings", "sectPrChange"],
}


def _local(element) -> str:
    return element.tag.rsplit("}", 1)[-1]


def _insert(parent, child) -> None:
    """Chèn ``child`` vào ``parent`` đúng thứ tự lược đồ; thẻ cùng tên đã có thì thay."""
    order = _ORDER[_local(parent)]
    name = _local(child)
    rank = order.index(name)
    for existing in list(parent):
        if _local(existing) == name and name not in {"headerReference", "footerReference"}:
            parent.replace(existing, child)
            return
    for existing in parent:
        local = _local(existing)
        if local in order and order.index(local) > rank:
            existing.addprevious(child)
            return
    parent.append(child)


def _xml(fragment: str):
    """Phần tử WordprocessingML từ chuỗi XML ngắn, tự thêm khai báo không gian tên w vào thẻ gốc."""
    return parse_xml(re.sub(r"^<(w:[A-Za-z]+)", lambda match: f"<{match.group(1)} {nsdecls('w')}", fragment, count=1))


def _border(side: str, value: str = "single", size: int = 6, color: str = "000000", space: int = 0) -> str:
    return f'<w:{side} w:val="{value}" w:sz="{size}" w:space="{space}" w:color="{color}"/>'


def _no_border(side: str) -> str:
    return f'<w:{side} w:val="nil"/>'


def _shade(fill: str):
    return _xml(f'<w:shd w:val="clear" w:color="auto" w:fill="{fill}"/>')


def _twips(cm: float) -> int:
    return int(round(cm * TWIPS))


def render(title: str, content: str, layout: str = "classic", images: dict | None = None,
           toc_pages: list | None = None) -> bytes:
    cover, blocks = base.parse_document(content)
    writer = _Writer(title, theme_for(layout), images or {}, toc_pages or [])
    shown, blocks = base.body_blocks(title, blocks)
    writer.shown = shown
    return writer.build(cover, blocks)


class _Writer:
    def __init__(self, title: str, theme: Theme, images: dict, toc_pages: list):
        self.title = " ".join(title.split())
        self.shown = self.title
        self.theme = theme
        self.images = images
        self.toc_pages = toc_pages
        self.document = Document()
        self.figures = 0
        self.tables = 0
        self.num_ids: dict[int, int] = {}
        self.drawing_id = 9000
        self.numbering = self.document.part.numbering_part.element
        self.next_abstract = max((int(item.get(qn("w:abstractNumId"))) for item in
                                  self.numbering.findall(qn("w:abstractNum"))), default=0) + 1

    # ---------- khung ----------
    def build(self, cover: Cover | None, blocks: list) -> bytes:
        self.page_setup()
        self.style_setup()
        if cover is not None:
            self.cover(cover)
        else:
            self.title_block()
        self.header_footer(cover)
        self.body(blocks)
        self.document.core_properties.title = self.title
        self.document.core_properties.author = "Peto"
        output = BytesIO()
        self.document.save(output)
        return output.getvalue()

    def page_setup(self) -> None:
        section = self.document.sections[0]
        section.page_width, section.page_height = Cm(21), Cm(29.7)
        section.top_margin, section.bottom_margin = Cm(2), Cm(2)
        section.left_margin, section.right_margin = Cm(3), Cm(2)
        section.header_distance, section.footer_distance = Cm(1), Cm(1)

    def _font(self, style, name: str = FONT) -> None:
        style.font.name = name
        fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        for attribute in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            fonts.attrib.pop(qn(attribute), None)
        fonts.set(qn("w:eastAsia"), name)
        fonts.set(qn("w:cs"), name)

    def style_setup(self) -> None:
        theme = self.theme
        styles = self.document.styles
        normal = styles["Normal"]
        self._font(normal)
        normal.font.size = Pt(13)
        normal.font.color.rgb = RGBColor(0, 0, 0)
        layout = normal.paragraph_format
        layout.line_spacing = 1.5
        layout.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        layout.first_line_indent = Cm(1)
        layout.space_before, layout.space_after = Pt(0), Pt(6)
        # Đề mục: phông theme của mẫu python-docx thắng tên phông đặt vào (Word hiện Calibri), nên bỏ thuộc tính theme.
        spacing = [(12, 6), (10, 4), (8, 3), (6, 3)]
        for level in range(1, 5):
            heading = styles[f"Heading {level}"]
            self._font(heading)
            heading.font.size = Pt(theme.heading_sizes[level - 1])
            heading.font.bold = True
            heading.font.italic = theme.heading_italic[level - 1]
            heading.font.all_caps = theme.heading_caps[level - 1]
            heading.font.color.rgb = RGBColor.from_string(theme.heading_color)
            heading_layout = heading.paragraph_format
            heading_layout.alignment = WD_ALIGN_PARAGRAPH.LEFT
            heading_layout.first_line_indent = Cm(0)
            heading_layout.space_before, heading_layout.space_after = Pt(spacing[level - 1][0]), Pt(spacing[level - 1][1])
            heading_layout.line_spacing = 1.3
            heading_layout.keep_with_next = True
            border = heading.element.get_or_add_pPr().find(qn("w:pBdr"))
            if border is not None:
                heading.element.pPr.remove(border)
            if level == 1 and theme.heading_rule:
                _insert(heading.element.get_or_add_pPr(),
                        _xml(f"<w:pBdr>{_border('bottom', size=8, color=theme.accent, space=3)}</w:pBdr>"))
        title = styles["Title"]
        self._font(title)
        title.font.color.rgb = RGBColor(0, 0, 0)
        border = title.element.get_or_add_pPr().find(qn("w:pBdr"))
        if border is not None:
            title.element.pPr.remove(border)
        caption = styles["Caption"]
        self._font(caption)
        caption.font.size = Pt(11.5 if theme.caption_small_caps else 12)
        caption.font.bold = False
        caption.font.italic = theme.name in {"classic", "essay"}
        caption.font.color.rgb = RGBColor(0, 0, 0)
        caption.paragraph_format.alignment = (WD_ALIGN_PARAGRAPH.LEFT if theme.caption_align == "left"
                                              else WD_ALIGN_PARAGRAPH.CENTER)
        caption.paragraph_format.first_line_indent = Cm(0)
        caption.paragraph_format.left_indent = Cm(1) if theme.caption_align == "left" else Cm(0)
        caption.paragraph_format.space_before, caption.paragraph_format.space_after = Pt(3), Pt(10)
        caption.paragraph_format.line_spacing = 1.15
        self._code_style()
        self._toc_styles()

    def _code_style(self) -> None:
        from docx.enum.style import WD_STYLE_TYPE
        theme = self.theme
        style = self.document.styles.add_style("Code", WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = self.document.styles["Normal"]
        self._font(style, theme.code_font)
        style.font.size = Pt(10.5 if theme.code == "bar" else 11)
        layout = style.paragraph_format
        layout.line_spacing = 1.0
        layout.alignment = WD_ALIGN_PARAGRAPH.LEFT
        layout.first_line_indent = Cm(0)
        layout.left_indent, layout.right_indent = Cm(0.2), Cm(0.2)
        layout.space_before, layout.space_after = Pt(0), Pt(0)
        layout.keep_together = True
        properties = style.element.get_or_add_pPr()
        # Đoạn liền nhau cùng viền được Word gom thành một khung: viền trên ở đoạn đầu, viền dưới ở đoạn cuối.
        if theme.code == "box":
            sides = "".join(_border(side, size=6, color="BFBFBF", space=4) for side in ("top", "left", "bottom", "right"))
            _insert(properties, _xml(f"<w:pBdr>{sides}</w:pBdr>"))
            _insert(properties, _shade("F2F2F2"))
        elif theme.code == "bar":
            _insert(properties, _xml(f"<w:pBdr>{_border('left', size=24, color=theme.accent, space=6)}</w:pBdr>"))
            _insert(properties, _shade("F3F6FA"))
        else:
            sides = "".join(_border(side, size=6, color="222222", space=4) for side in ("top", "bottom"))
            _insert(properties, _xml(f"<w:pBdr>{sides}</w:pBdr>"))

    def _toc_styles(self) -> None:
        """Kiểu "toc 1–3" của Word, có tab phải với dòng chấm dẫn tới số trang; Word dùng chúng khi cập nhật mục lục."""
        styles = self.document.styles.element
        right = _twips(TEXT_WIDTH)
        for level in (1, 2, 3):
            caps = "<w:caps/>" if level == 1 and self.theme.name == "classic" else ""
            bold = "<w:b/>" if level == 1 else ""
            color = f'<w:color w:val="{self.theme.accent}"/>' if level == 1 and self.theme.name == "band" else ""
            styles.append(parse_xml(
                f'<w:style {nsdecls("w")} w:type="paragraph" w:styleId="TOC{level}"><w:name w:val="toc {level}"/>'
                f'<w:basedOn w:val="Normal"/><w:next w:val="Normal"/><w:uiPriority w:val="39"/><w:unhideWhenUsed/>'
                f'<w:pPr><w:tabs><w:tab w:val="right" w:leader="dot" w:pos="{right}"/></w:tabs>'
                f'<w:spacing w:before="{60 if level == 1 else 0}" w:after="40" w:line="300" w:lineRule="auto"/>'
                f'<w:ind w:left="{(level - 1) * 360}" w:firstLine="0"/><w:jc w:val="left"/></w:pPr>'
                f"<w:rPr>{bold}{caps}{color}</w:rPr></w:style>"))

    # ---------- chữ ----------
    def run(self, paragraph, text: str, *, size: float | None = None, bold: bool | None = None,
            italic: bool | None = None, caps: bool = False, small_caps: bool = False, color: str | None = None,
            underline: bool = False, letter_spacing: float = 0.0, font: str | None = None):
        run = paragraph.add_run(text)
        run.font.size = Pt(size) if size else None
        if bold is not None:
            run.font.bold = bold
        if italic is not None:
            run.font.italic = italic
        if caps:
            run.font.all_caps = True
        if small_caps:
            run.font.small_caps = True
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
        if underline:
            run.font.underline = True
        if font:
            run.font.name = font
            run._r.get_or_add_rPr().get_or_add_rFonts().set(qn("w:cs"), font)
        if letter_spacing:
            _insert(run._r.get_or_add_rPr(), _xml(f'<w:spacing w:val="{int(letter_spacing * 20)}"/>'))
        return run

    def line(self, container, text: str = "", *, align=WD_ALIGN_PARAGRAPH.CENTER, before: float = 0, after: float = 0,
             spacing: float = 1.15, left: float | None = None, first: bool = False, **style):
        """Một đoạn định dạng sẵn (bìa, đầu trang). ``first``: dùng đoạn trống sẵn có của ô bảng."""
        paragraph = container.paragraphs[0] if first and container.paragraphs and not container.paragraphs[0].text \
            else container.add_paragraph()
        layout = paragraph.paragraph_format
        layout.first_line_indent = Cm(0)
        layout.space_before, layout.space_after = Pt(before), Pt(after)
        layout.line_spacing = spacing
        paragraph.alignment = align
        if left is not None:
            layout.left_indent = Cm(left)
        if text:
            self.run(paragraph, text, **style)
        return paragraph

    def write(self, paragraph, runs: list) -> None:
        theme = self.theme
        for item in runs:
            if item.math:
                paragraph._p.append(base.omml(item.text, False, FONT))
                continue
            run = paragraph.add_run(item.text)
            run.bold = item.bold or None
            run.italic, run.font.strike = item.italic or None, item.strike or None
            if item.code:
                run.font.name = theme.code_font
                run.font.size = Pt(11)
                if theme.code == "bar":
                    _insert(run._r.get_or_add_rPr(), _shade("E9EEF5"))

    def field(self, paragraph, kind: str, dirty: bool = False) -> None:
        char = OxmlElement("w:fldChar")
        char.set(qn("w:fldCharType"), kind)
        if dirty:
            char.set(qn("w:dirty"), "true")
        paragraph.add_run()._r.append(char)

    # ---------- bảng ẩn viền dùng cho bìa ----------
    def layout_table(self, heights: list[float], container=None):
        container = container or self.document
        table = container.add_table(rows=len(heights), cols=1)
        properties = table._tbl.tblPr
        _insert(properties, _xml(f'<w:tblW w:w="{_twips(TEXT_WIDTH)}" w:type="dxa"/>'))
        _insert(properties, _xml("<w:tblBorders>" + "".join(_no_border(side) for side in
                                 ("top", "left", "bottom", "right", "insideH", "insideV")) + "</w:tblBorders>"))
        _insert(properties, _xml('<w:tblLayout w:type="fixed"/>'))
        _insert(properties, _xml('<w:tblCellMar><w:left w:w="0" w:type="dxa"/><w:right w:w="0" w:type="dxa"/>'
                                 "</w:tblCellMar>"))
        cells = []
        for row, height in zip(table.rows, heights):
            _insert(row._tr.get_or_add_trPr(), _xml(f'<w:trHeight w:val="{_twips(height)}" w:hRule="atLeast"/>'))
            cell = row.cells[0]
            cell.width = Cm(TEXT_WIDTH)
            cells.append(cell)
        cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        if len(cells) == 3:
            cells[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cells[-1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
        return table, cells

    def end_cover(self) -> None:
        """Ngắt sang trang sau bìa bằng một đoạn rất thấp, để bìa cao gần hết trang mà không đẩy sang trang thứ hai."""
        paragraph = self.document.add_paragraph()
        layout = paragraph.paragraph_format
        layout.space_before = layout.space_after = Pt(0)
        layout.line_spacing = Pt(2)
        layout.first_line_indent = Cm(0)
        run = paragraph.add_run()
        run.font.size = Pt(2)
        run.add_break(WD_BREAK.PAGE)

    def members_table(self, cell, cover: Cover, style: str, indent: float = 0.0):
        members = cover.members
        has_code = any(member.code for member in members)
        has_note = any(member.note for member in members)
        headers = ["STT", "Họ và tên"] + (["MSSV"] if has_code else []) + (["Ghi chú"] if has_note else [])
        widths = [1.5, 6.0] + ([3.0] if has_code else []) + ([3.2] if has_note else [])
        table = cell.add_table(rows=1 + len(members), cols=len(headers))
        properties = table._tbl.tblPr
        total = sum(widths)
        _insert(properties, _xml(f'<w:tblW w:w="{_twips(total)}" w:type="dxa"/>'))
        _insert(properties, _xml('<w:tblLayout w:type="fixed"/>'))
        if indent:
            _insert(properties, _xml(f'<w:tblInd w:w="{_twips(indent)}" w:type="dxa"/>'))
        if style == "center":
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
        borders = {
            "grid": "".join(_border(side, size=4) for side in ("top", "left", "bottom", "right", "insideH", "insideV")),
            "band": "".join([_no_border("top"), _no_border("left"), _border("bottom", size=6, color="D5DDE8"),
                             _no_border("right"), _border("insideH", size=6, color="D5DDE8"), _no_border("insideV")]),
            "booktabs": "".join([_border("top", size=10), _no_border("left"), _border("bottom", size=10),
                                 _no_border("right"), _no_border("insideH"), _no_border("insideV")]),
        }
        kind = "grid" if self.theme.cover == "frame" else "band" if self.theme.cover == "band" else "booktabs"
        _insert(properties, _xml(f"<w:tblBorders>{borders[kind]}</w:tblBorders>"))
        rows = [headers] + [[str(index + 1), member.name] + ([member.code] if has_code else []) +
                            ([member.note] if has_note else []) for index, member in enumerate(members)]
        for row_index, (row, values) in enumerate(zip(table.rows, rows)):
            for column, (target, value) in enumerate(zip(row.cells, values)):
                target.width = Cm(widths[column])
                paragraph = target.paragraphs[0]
                paragraph.paragraph_format.first_line_indent = Cm(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.15
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if column == 0 else WD_ALIGN_PARAGRAPH.LEFT
                header = row_index == 0
                color = self.theme.accent if header and kind == "band" else None
                self.run(paragraph, value, size=12 if kind != "band" else 11.5, bold=header or None, color=color)
                if header and kind == "grid":
                    _insert(target._tc.get_or_add_tcPr(), _shade("F2F2F2"))
                if header and kind == "band":
                    _insert(target._tc.get_or_add_tcPr(), _shade("E9EEF5"))
                if header and kind == "booktabs":
                    _insert(target._tc.get_or_add_tcPr(), _xml(f"<w:tcBorders>{_border('bottom', size=5)}</w:tcBorders>"))
        return table

    # ---------- bìa ----------
    def cover(self, cover: Cover) -> None:
        section = self.document.sections[0]
        section.different_first_page_header_footer = True
        getattr(self, f"cover_{self.theme.cover}")(cover)
        self.end_cover()

    def cover_frame(self, cover: Cover) -> None:
        sectpr = self.document.sections[0]._sectPr
        sides = "".join(_border(side, "thickThinSmallGap", 24, "1A1A1A", 24) for side in ("top", "left", "bottom", "right"))
        _insert(sectpr, _xml(f'<w:pgBorders w:offsetFrom="page" w:display="firstPage">{sides}</w:pgBorders>'))
        _, (top, middle, bottom) = self.layout_table([5.4, 8.6, 10.4])
        first = True
        if cover.authority:
            self.line(top, cover.authority, size=13, caps=True, first=True)
            first = False
        self.line(top, cover.school or f"TRƯỜNG {DOTS}", size=14, bold=True, caps=True, first=first, before=2)
        self.line(top, cover.faculty or f"KHOA {DOTS}", size=13, bold=True, caps=True)
        self.line(top, "———  ◆  ———", size=12, before=4, after=10)
        self.logo(top, cover)
        topic = cover.topic or self.title
        self.line(middle, cover.kind or "Báo cáo", size=21, bold=True, caps=True, first=True, after=4)
        if cover.course:
            self.line(middle, f"Môn: {cover.course}", size=15, bold=True, caps=True, after=18)
        if cover.topic:
            self.line(middle, "Đề tài", size=13, bold=True, caps=True, underline=True, before=6, after=4)
        self.line(middle, topic, size=17, bold=True, caps=True, after=2)
        if cover.subtitle:
            self.line(middle, f"({cover.subtitle})", size=13, italic=True)
        self.info_lines(bottom, cover, left=2.0, first=True)
        self.line(bottom, cover.when(), size=13, bold=True, before=22)

    def cover_band(self, cover: Cover) -> None:
        _, (top, middle, bottom) = self.layout_table([3.2, 15.6, 5.6])
        paragraph = self.line(top, cover.school or f"Trường {DOTS}", align=WD_ALIGN_PARAGRAPH.LEFT, size=10.5,
                              bold=True, caps=True, color=self.theme.accent, letter_spacing=1.4, first=True)
        paragraph._p.insert(1, self.band_shape())
        self.line(top, cover.faculty or f"Khoa {DOTS}", align=WD_ALIGN_PARAGRAPH.LEFT, size=10.5, caps=True,
                  color=self.theme.accent, letter_spacing=1.0, before=3)
        self.logo(top, cover, align=WD_ALIGN_PARAGRAPH.LEFT)
        self.line(middle, cover.kind or "Báo cáo", align=WD_ALIGN_PARAGRAPH.LEFT, size=12, bold=True, caps=True,
                  color=self.theme.accent, letter_spacing=1.6, first=True)
        if cover.course:
            self.line(middle, f"Môn {cover.course}", align=WD_ALIGN_PARAGRAPH.LEFT, size=14, color="44546A", before=3)
        self.line(middle, cover.topic or self.title, align=WD_ALIGN_PARAGRAPH.LEFT, size=30, bold=True,
                  color=self.theme.accent, before=14, spacing=1.05)
        if cover.subtitle:
            self.line(middle, cover.subtitle, align=WD_ALIGN_PARAGRAPH.LEFT, size=14, color="44546A", before=8)
        rule = self.line(middle, align=WD_ALIGN_PARAGRAPH.LEFT, before=4, after=18)
        rule.paragraph_format.right_indent = Cm(TEXT_WIDTH - 5.3)
        _insert(rule._p.get_or_add_pPr(), _xml(f"<w:pBdr>{_border('bottom', size=24, color=self.theme.accent)}</w:pBdr>"))
        pairs = [("Giảng viên hướng dẫn", cover.instructor or DOTS)]
        if cover.group:
            pairs.append(("Nhóm thực hiện", cover.group))
        if cover.class_name:
            pairs.append(("Lớp", cover.class_name))
        if not cover.members and not cover.group:
            pairs.append(("Sinh viên thực hiện", DOTS))
        table = middle.add_table(rows=len(pairs), cols=2)
        properties = table._tbl.tblPr
        _insert(properties, _xml("<w:tblBorders>" + "".join(_no_border(side) for side in
                                 ("top", "left", "bottom", "right", "insideH", "insideV")) + "</w:tblBorders>"))
        _insert(properties, _xml('<w:tblLayout w:type="fixed"/>'))
        for row, (label, value) in zip(table.rows, pairs):
            for target, text, width, style in ((row.cells[0], label, 5.0, {"color": GRAY}),
                                               (row.cells[1], value, 10.0, {"bold": True})):
                target.width = Cm(width)
                paragraph = target.paragraphs[0]
                paragraph.paragraph_format.first_line_indent = Cm(0)
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.line_spacing = 1.15
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                self.run(paragraph, text, size=12.5, **style)
        if cover.members:
            self.line(middle, after=4)
            self.members_table(middle, cover, "left")
        self.line(bottom, cover.when(), align=WD_ALIGN_PARAGRAPH.LEFT, size=12, color=GRAY, first=True)

    def cover_rules(self, cover: Cover) -> None:
        _, (top, middle, bottom) = self.layout_table([4.4, 9.4, 10.6])
        self.line(top, cover.school or f"Trường {DOTS}", size=14, small_caps=True, letter_spacing=1.4, first=True)
        self.line(top, cover.faculty or f"Khoa {DOTS}", size=13, italic=True, before=3)
        self.logo(top, cover)
        kind = cover.kind or "Báo cáo"
        label = f"{kind} · {cover.course}" if cover.course else kind
        opening = self.line(middle, label, size=13, small_caps=True, letter_spacing=1.0, first=True, before=16)
        _insert(opening._p.get_or_add_pPr(), _xml(f"<w:pBdr>{_border('top', size=6, color='111111', space=16)}</w:pBdr>"))
        closing = self.line(middle, cover.topic or self.title, size=25, before=8, spacing=1.1,
                            after=0 if cover.subtitle else 16)
        if cover.subtitle:
            closing = self.line(middle, cover.subtitle, size=13, italic=True, before=8, after=16)
        _insert(closing._p.get_or_add_pPr(), _xml(f"<w:pBdr>{_border('bottom', size=6, color='111111', space=16)}</w:pBdr>"))
        self.line(bottom, f"Giảng viên hướng dẫn: {cover.instructor or DOTS}", size=13, first=True)
        for extra in (cover.group, cover.class_name and f"Lớp {cover.class_name}"):
            if extra:
                self.line(bottom, extra, size=13)
        if cover.members:
            self.line(bottom, after=6)
            self.members_table(bottom, cover, "center")
        elif not cover.group:
            self.line(bottom, f"Sinh viên thực hiện: {DOTS}", size=13)
        self.line(bottom, cover.when(), size=13, small_caps=True, letter_spacing=1.0, before=26)

    def info_lines(self, cell, cover: Cover, left: float, first: bool) -> None:
        paragraph = self.line(cell, align=WD_ALIGN_PARAGRAPH.LEFT, left=left, first=first, after=2)
        self.run(paragraph, "Giảng viên hướng dẫn: ", size=13, bold=True)
        self.run(paragraph, cover.instructor or DOTS, size=13)
        for label, value in (("Nhóm thực hiện: ", cover.group), ("Lớp: ", cover.class_name)):
            if value:
                paragraph = self.line(cell, align=WD_ALIGN_PARAGRAPH.LEFT, left=left, after=2)
                self.run(paragraph, label, size=13, bold=True)
                self.run(paragraph, value, size=13)
        if cover.members:
            self.line(cell, after=2)
            self.members_table(cell, cover, "left", indent=left)
        elif not cover.group:
            paragraph = self.line(cell, align=WD_ALIGN_PARAGRAPH.LEFT, left=left)
            self.run(paragraph, "Sinh viên thực hiện: ", size=13, bold=True)
            self.run(paragraph, DOTS, size=13)

    def logo(self, cell, cover: Cover, align=WD_ALIGN_PARAGRAPH.CENTER) -> None:
        image = self.images.get(cover.logo) if cover.logo else None
        if image is None:
            return
        width, height = base.fit(image, 200, 92)
        paragraph = self.line(cell, align=align, before=4, after=4, spacing=1.0)
        paragraph.add_run().add_picture(BytesIO(image.data), height=Pt(height))

    def band_shape(self):
        """Dải xanh sát mép trái trang bìa: hình chữ nhật neo theo trang, nằm sau chữ (kèm bản VML cho Word cũ)."""
        self.drawing_id += 1
        width, height = 508000, 10692000  # EMU: 40 pt × cả chiều cao A4
        return parse_xml(
            f'<w:r {nsdecls("w")} xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
            'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
            'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
            'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
            'xmlns:v="urn:schemas-microsoft-com:vml"><mc:AlternateContent><mc:Choice Requires="wps"><w:drawing>'
            '<wp:anchor distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="251658240" behindDoc="1" '
            'locked="1" layoutInCell="0" allowOverlap="1"><wp:simplePos x="0" y="0"/>'
            '<wp:positionH relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionH>'
            '<wp:positionV relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionV>'
            f'<wp:extent cx="{width}" cy="{height}"/><wp:effectExtent l="0" t="0" r="0" b="0"/><wp:wrapNone/>'
            f'<wp:docPr id="{self.drawing_id}" name="Dải màu bìa"/><wp:cNvGraphicFramePr/><a:graphic>'
            '<a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"><wps:wsp>'
            f'<wps:cNvSpPr/><wps:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="{width}" cy="{height}"/></a:xfrm>'
            '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom>'
            f'<a:solidFill><a:srgbClr val="{self.theme.accent}"/></a:solidFill><a:ln><a:noFill/></a:ln></wps:spPr>'
            '<wps:bodyPr/></wps:wsp></a:graphicData></a:graphic></wp:anchor></w:drawing></mc:Choice>'
            '<mc:Fallback><w:pict><v:rect style="position:absolute;margin-left:0;margin-top:0;width:40pt;height:842pt;'
            'z-index:-251658240;mso-position-horizontal-relative:page;mso-position-vertical-relative:page" '
            f'fillcolor="#{self.theme.accent}" stroked="f"/></w:pict></mc:Fallback></mc:AlternateContent></w:r>')

    # ---------- tên tài liệu khi không có bìa ----------
    def title_block(self) -> None:
        theme = self.theme
        paragraph = self.document.add_paragraph(style="Title")
        layout = paragraph.paragraph_format
        layout.first_line_indent = Cm(0)
        layout.space_before = Pt(0)
        if theme.title == "center":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            layout.space_after = Pt(14)
            self.run(paragraph, self.shown, size=15, bold=True, caps=True)
        elif theme.title == "band":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            layout.space_after = Pt(14)
            self.run(paragraph, self.shown, size=20, bold=True, color=theme.accent)
            _insert(paragraph._p.get_or_add_pPr(), _xml(f"<w:pBdr>{_border('bottom', size=8, color=theme.accent, space=6)}</w:pBdr>"))
        elif theme.title == "rule":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            layout.space_after = Pt(14)
            self.run(paragraph, self.shown, size=18)
            _insert(paragraph._p.get_or_add_pPr(), _xml(f"<w:pBdr>{_border('bottom', size=6, color='111111', space=6)}</w:pBdr>"))
        else:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            layout.space_after = Pt(18)
            self.run(paragraph, self.shown, size=18, bold=True)

    # ---------- đầu, chân trang ----------
    def header_footer(self, cover: Cover | None) -> None:
        theme = self.theme
        section = self.document.sections[0]
        if cover is None and theme.header:
            section.different_first_page_header_footer = True
        if theme.header:
            paragraph = section.header.paragraphs[0]
            layout = paragraph.paragraph_format
            layout.first_line_indent = Cm(0)
            layout.space_after = Pt(0)
            layout.line_spacing = 1.0
            style = {"size": 10, "italic": True} if theme.header == "title" else \
                {"size": 10, "small_caps": theme.header_small_caps, "color": GRAY if theme.name == "band" else "333333"}
            if theme.header == "title":
                paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                self.run(paragraph, base.running_title(self.title), **style)
            else:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                left = ""
                if cover is not None:
                    left = " · ".join(part for part in (cover.course, cover.group) if part) or cover.kind
                right = base.running_title(self.title) if left else ""
                left = left or base.running_title(self.title)
                _insert(paragraph._p.get_or_add_pPr(),
                        _xml(f'<w:tabs><w:tab w:val="right" w:pos="{_twips(TEXT_WIDTH)}"/></w:tabs>'))
                self.run(paragraph, left, **style)
                if right:
                    self.run(paragraph, "\t" + right, **style)
                color = "C9D3E0" if theme.name == "band" else "555555"
                _insert(paragraph._p.get_or_add_pPr(),
                        _xml(f"<w:pBdr>{_border('bottom', size=4, color=color, space=3)}</w:pBdr>"))
        footers = [section.footer] + ([section.first_page_footer] if cover is None and theme.header else [])
        for footer in footers:
            paragraph = footer.paragraphs[0]
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.space_after = Pt(0)
            if theme.footer == "number":
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                self.page_number(paragraph, "", 12, None)
            elif theme.footer == "trang":
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                self.page_number(paragraph, "Trang ", 11, None)
            else:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                self.page_number(paragraph, "Trang ", 10.5, GRAY)

    def page_number(self, paragraph, prefix: str, size: float, color: str | None) -> None:
        if prefix:
            self.run(paragraph, prefix, size=size, color=color)
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        run = OxmlElement("w:r")
        properties = OxmlElement("w:rPr")
        size_element = OxmlElement("w:sz")
        size_element.set(qn("w:val"), str(int(size * 2)))
        if color:
            color_element = OxmlElement("w:color")
            color_element.set(qn("w:val"), color)
            properties.append(color_element)
        properties.append(size_element)
        run.append(properties)
        text = OxmlElement("w:t")
        text.text = "1"
        run.append(text)
        field.append(run)
        paragraph._p.append(field)

    # ---------- phần thân ----------
    def body(self, blocks: list) -> None:
        index = 0
        entries = base.contents(blocks)
        while index < len(blocks):
            block = blocks[index]
            following = blocks[index + 1] if index + 1 < len(blocks) else None
            caption = base.table_caption(block, following)
            if caption is not None:
                self.table(following.rows, caption)
                index += 2
                continue
            if block.kind == "code":
                end = index
                while end < len(blocks) and blocks[end].kind == "code":
                    end += 1
                self.code([item.runs for item in blocks[index:end]])
                index = end
                continue
            self.block(block, entries)
            index += 1

    def block(self, block, entries: list) -> None:
        document = self.document
        if block.kind == "table" and block.rows:
            self.table(block.rows, None)
        elif block.kind == "rule":
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.first_line_indent = Cm(0)
            _insert(paragraph._p.get_or_add_pPr(), _xml(f"<w:pBdr>{_border('bottom', size=6, color='BFBFBF', space=1)}</w:pBdr>"))
        elif block.kind == "math":
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph._p.append(base.omml(block.math, True, FONT))
        elif block.kind == "image":
            self.image(block)
        elif block.kind == "slot":
            self.slot(block.caption)
        elif block.kind == "toc":
            self.toc(entries)
        elif block.kind == "heading":
            paragraph = document.add_paragraph(style=f"Heading {block.level}")
            self.write(paragraph, block.runs)
        else:
            paragraph = document.add_paragraph()
            layout = paragraph.paragraph_format
            if block.list_info:
                number_x, text_x = base.list_indent(block.list_info.depth)
                layout.left_indent = Pt(text_x)
                layout.first_line_indent = Pt(number_x - text_x) if block.number is not None else Pt(0)
                if block.number is not None:
                    numbering = paragraph._p.get_or_add_pPr().get_or_add_numPr()
                    numbering.get_or_add_ilvl().val = block.list_info.depth
                    numbering.get_or_add_numId().val = self.num_id(block.list_info)
                layout.space_after = Pt(3)
            if block.kind == "quote":
                layout.left_indent = Cm(1)
                layout.first_line_indent = Cm(0)
                if self.theme.name == "band":
                    _insert(paragraph._p.get_or_add_pPr(),
                            _xml(f"<w:pBdr>{_border('left', size=18, color=self.theme.accent, space=8)}</w:pBdr>"))
                for item in block.runs:
                    item.italic = True
            self.write(paragraph, block.runs)

    def num_id(self, info) -> int:
        """Mỗi danh sách một định nghĩa đánh số riêng: số tự bắt đầu lại, Word đánh lại số khi thêm bớt mục."""
        if info.id not in self.num_ids:
            levels = "".join(base.numbering_level(level, info) for level in range(9))
            abstract = parse_xml(f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{self.next_abstract}">'
                                 f'<w:multiLevelType w:val="hybridMultilevel"/>{levels}</w:abstractNum>')
            first = self.numbering.find(qn("w:num"))
            if first is None:
                self.numbering.append(abstract)
            else:
                first.addprevious(abstract)
            self.num_ids[info.id] = self.numbering.add_num(self.next_abstract).numId
            self.next_abstract += 1
        return self.num_ids[info.id]

    def code(self, lines: list) -> None:
        paragraphs = []
        for runs in lines:
            paragraph = self.document.add_paragraph(style="Code")
            for item in runs:
                paragraph.add_run(item.text)
            paragraphs.append(paragraph)
        paragraphs[0].paragraph_format.space_before = Pt(4)
        paragraphs[-1].paragraph_format.space_after = Pt(8)

    def caption(self, kind: str, number: int, text: str):
        theme = self.theme
        paragraph = self.document.add_paragraph(style="Caption")
        label = f"{kind} {number}{theme.caption_separator}"
        self.run(paragraph, label, bold=True, italic=False if theme.name in {"classic", "essay"} else None,
                 small_caps=theme.caption_small_caps, color=theme.caption_color if theme.caption_color != "000000" else None)
        if text:
            self.run(paragraph, " " + text)
        return paragraph

    def image(self, block) -> None:
        image = self.images.get(block.image)
        if image is None:
            paragraph = self.document.add_paragraph()
            paragraph.add_run(base.placeholder(block))
            return
        width, _ = base.fit(image, TEXT_WIDTH * 28.35, 25.7 * 28.35 * .55)
        picture = self.document.add_paragraph()
        picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
        layout = picture.paragraph_format
        layout.first_line_indent = Cm(0)
        layout.line_spacing = 1
        layout.space_before = Pt(4)
        layout.keep_with_next = bool(block.caption)
        picture.add_run().add_picture(BytesIO(image.data), width=Pt(width))
        if block.caption:
            self.figures += 1
            self.caption("Hình", self.figures, base.figure_caption(block.caption))

    def slot(self, caption: str) -> None:
        """Khung nét đứt chừa chỗ dán ảnh chụp màn hình, kèm chú thích "Hình N"."""
        table = self.document.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        properties = table._tbl.tblPr
        width = TEXT_WIDTH * 0.85
        _insert(properties, _xml(f'<w:tblW w:w="{_twips(width)}" w:type="dxa"/>'))
        _insert(properties, _xml('<w:tblLayout w:type="fixed"/>'))
        sides = "".join(_border(side, "dashed", 6, "9AA3AD") for side in ("top", "left", "bottom", "right"))
        _insert(properties, _xml(f"<w:tblBorders>{sides}</w:tblBorders>"))
        row = table.rows[0]
        _insert(row._tr.get_or_add_trPr(), _xml(f'<w:trHeight w:val="{_twips(5.4)}" w:hRule="exact"/>'))
        cell = row.cells[0]
        cell.width = Cm(width)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        _insert(cell._tc.get_or_add_tcPr(), _shade("FAFBFC"))
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.first_line_indent = Cm(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.2
        paragraph.paragraph_format.keep_with_next = True  # khung đi cùng chú thích "Hình N" bên dưới
        text = base.figure_caption(caption)
        self.run(paragraph, "Chỗ dán ảnh chụp màn hình" + (f":\n{text}" if text else ""), size=11, italic=True,
                 color="6B7480")
        self.figures += 1
        self.caption("Hình", self.figures, text).paragraph_format.space_before = Pt(6)

    def table(self, rows: list, caption: str | None) -> None:
        theme = self.theme
        if caption is not None:
            self.tables += 1
            paragraph = self.caption("Bảng", self.tables, caption)
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.space_before = Pt(6)
            paragraph.paragraph_format.space_after = Pt(4)
        columns = max(len(row) for row in rows)
        rows = [row + [[] for _ in range(columns - len(row))] for row in rows]
        widths = [width / 28.3465 for width in base.column_widths(rows, TEXT_WIDTH * 28.3465)]
        centered = base.centered_columns(rows)
        table = self.document.add_table(rows=0, cols=columns)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        properties = table._tbl.tblPr
        _insert(properties, _xml(f'<w:tblW w:w="{_twips(TEXT_WIDTH)}" w:type="dxa"/>'))
        _insert(properties, _xml('<w:tblLayout w:type="fixed"/>'))
        _insert(properties, _xml('<w:tblCellMar><w:top w:w="40" w:type="dxa"/><w:left w:w="100" w:type="dxa"/>'
                                 '<w:bottom w:w="40" w:type="dxa"/><w:right w:w="100" w:type="dxa"/></w:tblCellMar>'))
        borders = {
            "grid": "".join(_border(side, size=6) for side in ("top", "left", "bottom", "right", "insideH", "insideV")),
            "band": "".join([_no_border("top"), _no_border("left"), _border("bottom", size=6, color="C9D3E0"),
                             _no_border("right"), _border("insideH", size=6, color="C9D3E0"), _no_border("insideV")]),
            "booktabs": "".join([_border("top", size=12), _no_border("left"), _border("bottom", size=12),
                                 _no_border("right"), _no_border("insideH"), _no_border("insideV")]),
        }[theme.table]
        _insert(properties, _xml(f"<w:tblBorders>{borders}</w:tblBorders>"))
        grid = table._tbl.tblGrid
        for column, width in zip(grid.findall(qn("w:gridCol")), widths):
            column.set(qn("w:w"), str(_twips(width)))
        # Bảng ngắn không tách trang, như bản PDF: mọi hàng trừ hàng cuối "giữ với đoạn sau" và không hàng nào ngắt giữa.
        short = len(rows) <= base.SHORT_TABLE_ROWS
        for row_index, row in enumerate(rows):
            cells = table.add_row().cells
            header = row_index == 0
            if short:
                _insert(table.rows[row_index]._tr.get_or_add_trPr(), _xml("<w:cantSplit/>"))
            for column, (cell, runs) in enumerate(zip(cells, row)):
                cell.width = Cm(widths[column])
                paragraph = cell.paragraphs[0]
                layout = paragraph.paragraph_format
                layout.first_line_indent = Cm(0)
                layout.space_after = Pt(0)
                layout.line_spacing = 1.15
                if short and row_index < len(rows) - 1:
                    layout.keep_with_next = True
                paragraph.alignment = (WD_ALIGN_PARAGRAPH.CENTER if column in centered or (header and theme.table == "grid")
                                       else WD_ALIGN_PARAGRAPH.LEFT)
                self.write(paragraph, runs)
                for run in paragraph.runs:
                    run.font.size = Pt(12)
                    if header:
                        run.font.bold = True
                        if theme.table == "band":
                            run.font.color.rgb = RGBColor(255, 255, 255)
                cell_properties = cell._tc.get_or_add_tcPr()
                if header and theme.table == "grid":
                    _insert(cell_properties, _shade("F2F2F2"))
                elif header and theme.table == "band":
                    _insert(cell_properties, _shade(theme.accent))
                elif theme.table == "band" and row_index % 2 == 0:
                    _insert(cell_properties, _shade("F3F6FA"))
                if header and theme.table == "booktabs":
                    _insert(cell_properties, _xml(f"<w:tcBorders>{_border('bottom', size=6)}</w:tcBorders>"))
            if header:
                _insert(table.rows[0]._tr.get_or_add_trPr(), _xml("<w:tblHeader/>"))
        spacer = self.document.add_paragraph()
        spacer.paragraph_format.space_after = Pt(4)
        spacer.paragraph_format.line_spacing = 1.0

    def toc(self, entries: list) -> None:
        if not entries:
            return
        theme = self.theme
        document = self.document
        paragraph = document.add_paragraph()
        layout = paragraph.paragraph_format
        layout.first_line_indent = Cm(0)
        layout.space_after = Pt(14)
        layout.keep_with_next = True
        if theme.name == "band":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            self.run(paragraph, "Mục lục", size=20, bold=True, color=theme.accent)
        elif theme.name == "minimal":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            self.run(paragraph, "Mục lục", size=15, small_caps=True, letter_spacing=1.2)
        else:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            self.run(paragraph, "Mục lục", size=14 if theme.name == "classic" else 16, bold=True,
                     caps=theme.name == "classic")
        pages = {(level, text): page for level, text, page in self.toc_pages}
        last = None
        for position, entry in enumerate(entries):
            text = base.plain(entry.runs)
            item = document.add_paragraph(style=document.styles[f"toc {entry.level}"])
            item.paragraph_format.left_indent = Pt(18 * (entry.level - 1))
            item.paragraph_format.first_line_indent = Cm(0)
            if position == 0:
                # Trường mục lục đánh dấu cần cập nhật: Word hỏi cập nhật khi mở rồi dựng lại có số trang. Trình xem
                # khác (điện thoại, Google Docs) hiện danh sách viết sẵn, kèm số trang lấy từ bản PDF khi có.
                self.field(item, "begin", dirty=True)
                instruction = OxmlElement("w:instrText")
                instruction.set(qn("xml:space"), "preserve")
                instruction.text = f' TOC \\o "1-{base.TOC_LEVELS}" \\h \\z \\u '
                item.add_run()._r.append(instruction)
                self.field(item, "separate")
            item.add_run(text)
            page = pages.get((entry.level, text))
            if page:
                item.add_run(f"\t{page}")
            last = item
        self.field(last, "end")
        document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
